"""Download the pinned official UCI Online Retail workbook (dataset 352)."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import zipfile
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl import load_workbook


DATASET_ID = 352
SOURCE_URL = "https://archive.ics.uci.edu/static/public/352/online+retail.zip"
DATASET_PAGE = "https://archive.ics.uci.edu/dataset/352/online+retail"
DOI = "10.24432/C5BW33"
LICENSE = "CC BY 4.0"
ARCHIVE_SHA256 = "f5385cbb54bbebf7196389109c6b0621faab0c304e3702548165e71c84aede8b"
WORKBOOK_SHA256 = "43465a06f2ccf7c8b5bd2892bc7defb52f97487934fe93b16ae4c3936424676d"
ARCHIVE_MEMBER = "Online Retail.xlsx"
EXPECTED_ROWS = 541_909
EXPECTED_HEADERS = (
    "InvoiceNo",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "UnitPrice",
    "CustomerID",
    "Country",
)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def validate_workbook(path: Path, *, expected_rows: int = EXPECTED_ROWS) -> None:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Online Retail"]:
            raise ValueError("UCI Online Retail workbook sheet contract changed")
        sheet = workbook["Online Retail"]
        header = tuple(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        if header != EXPECTED_HEADERS:
            raise ValueError("UCI Online Retail workbook column contract changed")
        if sheet.max_row != expected_rows + 1:
            raise ValueError("UCI Online Retail workbook row count changed")
    finally:
        workbook.close()


def manifest_for(path: Path) -> dict[str, object]:
    return {
        "schema_version": 1,
        "dataset_id": DATASET_ID,
        "name": "Online Retail",
        "dataset_page": DATASET_PAGE,
        "doi": DOI,
        "license": LICENSE,
        "source_url": SOURCE_URL,
        "archive_member": ARCHIVE_MEMBER,
        "archive_sha256": ARCHIVE_SHA256,
        "workbook_sha256": sha256_file(path),
        "rows": EXPECTED_ROWS,
        "columns": list(EXPECTED_HEADERS),
    }


def download(output: Path) -> dict[str, object]:
    output.parent.mkdir(parents=True, exist_ok=True)
    archive_part = output.with_name(output.stem + ".part.zip")
    workbook_part = output.with_name(output.stem + ".part.xlsx")
    try:
        request = Request(SOURCE_URL, headers={"User-Agent": "ai-defense-lab/1.0"})
        with urlopen(request, timeout=120) as source, archive_part.open("wb") as target:
            shutil.copyfileobj(source, target)
        if sha256_file(archive_part) != ARCHIVE_SHA256:
            raise ValueError("official UCI archive SHA-256 differs from the pinned artifact")
        with zipfile.ZipFile(archive_part) as archive:
            if archive.namelist() != [ARCHIVE_MEMBER]:
                raise ValueError("official UCI archive member contract changed")
            with archive.open(ARCHIVE_MEMBER) as source, workbook_part.open("wb") as target:
                shutil.copyfileobj(source, target)
        if sha256_file(workbook_part) != WORKBOOK_SHA256:
            raise ValueError("official UCI workbook SHA-256 differs from the pinned artifact")
        validate_workbook(workbook_part)
        os.replace(workbook_part, output)
        manifest = manifest_for(output)
        output.with_suffix(".manifest.json").write_text(
            json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
        )
        return manifest
    finally:
        archive_part.unlink(missing_ok=True)
        workbook_part.unlink(missing_ok=True)


def verify_existing(output: Path) -> dict[str, object]:
    manifest_path = output.with_suffix(".manifest.json")
    if not output.is_file() or not manifest_path.is_file():
        raise FileNotFoundError("workbook and manifest must both exist")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest != manifest_for(output) or manifest["workbook_sha256"] != WORKBOOK_SHA256:
        raise ValueError("existing UCI workbook provenance does not match the pinned manifest")
    validate_workbook(output)
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("data/reference/online_retail.xlsx"))
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    manifest = download(args.output) if args.force or not args.output.exists() else verify_existing(args.output)
    print(json.dumps(manifest, sort_keys=True))


if __name__ == "__main__":
    main()
