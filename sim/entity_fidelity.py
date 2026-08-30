"""Entity-level fidelity audit against the real UCI Online Retail population."""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
import os
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook

from scripts.download_retail_reference import (
    DATASET_ID,
    EXPECTED_HEADERS,
    EXPECTED_ROWS,
    SOURCE_URL,
    WORKBOOK_SHA256,
    sha256_file,
)
from sim.population import (
    COPULA_PRIMITIVES,
    COPULA_QUANTILE_PROBABILITIES,
    SINGLETON_PRIMITIVES,
    BehaviorProfile,
    PopulationSimulator,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REFERENCE = ROOT / "data" / "reference" / "online_retail.xlsx"
DEFAULT_OUTPUT = ROOT / "results" / "entity_fidelity.json"
DEFAULT_PROTOCOL = ROOT / "results" / "entity_fidelity_v3_protocol.json"
DEFAULT_READINESS = ROOT / "results" / "entity_fidelity_v3_readiness.json"
DEFAULT_SEAL_RECEIPT = ROOT / "results" / "entity_fidelity_v3_sealed_receipt.json"
SEED = 20260827
CALIBRATION_FRACTION = 0.55
DIAGNOSTIC_FRACTION = 0.225
SEALED_FINAL_FRACTION = 0.225
SYNTHETIC_CUSTOMERS = 4_500
SYNTHETIC_EVENTS: int | None = None
BOOTSTRAP_REPEATS = 1_000
GATE_MAX_AUC_UPPER_95 = 0.65
READINESS_BOOTSTRAP_REPEATS = 10_000
READINESS_SEEDS = (20260827, 20260828, 20260829)
CROSS_FIT_FOLDS = 5
SEALED_FINAL_CONFIRMATION = "APPROVED_SINGLE_FINAL_SCORE"
COPULA_CORRELATION_SHRINKAGE = 0.10
EB_CIRCADIAN_PRIOR_STRENGTH = 4.0
EXPECTED_PARTITION_HASHES = {
    "calibration": "53afc6c77f75002c1359ef445e1d1a7527dfd3532c0c4f23f5e557e2b3ccd0de",
    "diagnostic": "b4f75eadaf75581dc46b725126364e30fc7df8cc7da616bd04b087e70036cb66",
    "sealed_final": "a0ec7b1f719bcbe8304eaf1748d040ada0a67daea4cad66648a40ce703e9cc2f",
}
FEATURES = (
    "median_log_amount_percentile",
    "amount_cv",
    "mean_hour_sin",
    "mean_hour_cos",
    "log1p_median_interpurchase_days",
    "log1p_frequency_per_30d",
    "log1p_recency_days",
    "log1p_max_burst_24h",
)


def _dependencies() -> tuple[Any, Any, Any, Any]:
    try:
        import lightgbm
        import numpy as np
        import pandas as pd
        from sklearn.metrics import roc_auc_score
    except ImportError as exc:
        raise RuntimeError(
            "entity fidelity requires numpy, pandas, scikit-learn, lightgbm, and openpyxl"
        ) from exc
    return np, pd, lightgbm, roc_auc_score


def _customer_id(value: Any) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def aggregate_retail_rows(rows: Iterable[Mapping[str, Any]]) -> tuple[Any, dict[str, int]]:
    """Filter line items and aggregate their value to one row per invoice."""
    np, pd, _, _ = _dependencies()
    invoices: dict[str, list[Any]] = {}
    counts = {
        "input_line_items": 0,
        "excluded_cancellation_line_items": 0,
        "excluded_missing_customer_line_items": 0,
        "excluded_nonpositive_quantity_line_items": 0,
        "excluded_nonpositive_price_line_items": 0,
        "usable_line_items": 0,
    }
    for row in rows:
        counts["input_line_items"] += 1
        invoice_id = str(row.get("InvoiceNo", "")).strip()
        customer_id = _customer_id(row.get("CustomerID"))
        try:
            quantity = float(row.get("Quantity"))
            unit_price = float(row.get("UnitPrice"))
        except (TypeError, ValueError):
            quantity = unit_price = float("nan")
        if invoice_id.upper().startswith("C"):
            counts["excluded_cancellation_line_items"] += 1
            continue
        if customer_id is None:
            counts["excluded_missing_customer_line_items"] += 1
            continue
        if not math.isfinite(quantity) or quantity <= 0:
            counts["excluded_nonpositive_quantity_line_items"] += 1
            continue
        if not math.isfinite(unit_price) or unit_price <= 0:
            counts["excluded_nonpositive_price_line_items"] += 1
            continue
        timestamp = pd.Timestamp(row.get("InvoiceDate"))
        if pd.isna(timestamp):
            raise ValueError(f"invoice {invoice_id!r} has an invalid InvoiceDate")
        if timestamp.tz is not None:
            timestamp = timestamp.tz_convert("UTC").tz_localize(None)
        existing = invoices.get(invoice_id)
        if existing is None:
            invoices[invoice_id] = [customer_id, timestamp, quantity * unit_price]
        else:
            if existing[0] != customer_id:
                raise ValueError(f"invoice {invoice_id!r} changes customer across line items")
            existing[1] = min(existing[1], timestamp)
            existing[2] += quantity * unit_price
        counts["usable_line_items"] += 1
    frame = pd.DataFrame(
        (
            {"customer_id": values[0], "ts": values[1], "amount": values[2]}
            for _, values in sorted(invoices.items())
        ),
        columns=["customer_id", "ts", "amount"],
    )
    if frame.empty or not np.isfinite(frame["amount"]).all() or (frame["amount"] <= 0).any():
        raise ValueError("no valid positive invoices remained after UCI filtering")
    counts["usable_invoices"] = len(frame)
    counts["usable_customers"] = int(frame["customer_id"].nunique())
    return frame.sort_values(["customer_id", "ts"], kind="stable").reset_index(drop=True), counts


def verify_reference_provenance(
    workbook: Path,
    manifest_path: Path | None = None,
    *,
    expected_sha256: str = WORKBOOK_SHA256,
) -> dict[str, Any]:
    manifest_path = manifest_path or workbook.with_suffix(".manifest.json")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    digest = sha256_file(workbook)
    required = {
        "dataset_id": DATASET_ID,
        "source_url": SOURCE_URL,
        "workbook_sha256": expected_sha256,
        "rows": EXPECTED_ROWS,
        "license": "CC BY 4.0",
    }
    if any(manifest.get(key) != value for key, value in required.items()) or digest != expected_sha256:
        raise ValueError("UCI Online Retail workbook or manifest provenance mismatch")
    return {
        **required,
        "path": str(workbook),
        "manifest_path": str(manifest_path),
        "manifest_sha256": sha256_file(manifest_path),
    }


def load_retail_transactions(workbook_path: Path) -> tuple[Any, dict[str, Any]]:
    provenance = verify_reference_provenance(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if workbook.sheetnames != ["Online Retail"]:
            raise ValueError("UCI Online Retail workbook sheet contract changed")
        values = workbook["Online Retail"].iter_rows(values_only=True)
        headers = tuple(next(values))
        if headers != EXPECTED_HEADERS:
            raise ValueError("UCI Online Retail workbook column contract changed")
        frame, counts = aggregate_retail_rows(dict(zip(headers, row, strict=True)) for row in values)
    finally:
        workbook.close()
    if counts["input_line_items"] != EXPECTED_ROWS:
        raise ValueError("UCI Online Retail workbook line-item count changed")
    return frame, {**provenance, **counts}


def generate_synthetic_transactions(
    *,
    events: int | None,
    customers: int,
    seed: int,
    behavior_profile: BehaviorProfile | None = None,
) -> Any:
    _, pd, _, _ = _dependencies()
    if events is None:
        if behavior_profile is None:
            raise ValueError("derived event count requires a behavior profile")
        events = round(behavior_profile.activity_mean_transactions * customers)
    simulator = PopulationSimulator(
        seed=seed,
        customers=customers,
        merchants_per_category=1,
        agentic_rate=0,
        behavior_profile=behavior_profile,
    )
    rows = (
        {
            "customer_id": event["txn"]["customer_id"],
            "ts": pd.Timestamp(event["ts"]).tz_convert("UTC").tz_localize(None),
            "amount": event["txn"]["amount_minor"] / 100,
        }
        for event in simulator.iter_legitimate(events)
    )
    return pd.DataFrame(rows).sort_values(["customer_id", "ts"], kind="stable").reset_index(drop=True)


def customer_split(
    customer_ids: Iterable[str],
    *,
    seed: int,
    namespace: str,
    calibration_fraction: float = CALIBRATION_FRACTION,
) -> tuple[set[str], set[str]]:
    if not 0 < calibration_fraction < 1:
        raise ValueError("calibration fraction must be strictly between zero and one")
    calibration: set[str] = set()
    evaluation: set[str] = set()
    cutoff = int(calibration_fraction * (1 << 64))
    for customer_id in sorted(set(customer_ids)):
        digest = hashlib.sha256(f"{seed}:{namespace}:{customer_id}".encode()).digest()
        (calibration if int.from_bytes(digest[:8], "big") < cutoff else evaluation).add(customer_id)
    if not calibration or not evaluation or calibration & evaluation:
        raise ValueError("customer split must be nonempty and disjoint")
    return calibration, evaluation


def customer_partitions(
    customer_ids: Iterable[str],
    *,
    seed: int,
    namespace: str,
) -> dict[str, set[str]]:
    all_ids = set(customer_ids)
    fractions = (CALIBRATION_FRACTION, DIAGNOSTIC_FRACTION, SEALED_FINAL_FRACTION)
    if not math.isclose(sum(fractions), 1.0, abs_tol=1e-12):
        raise ValueError("entity-fidelity partition fractions must sum to one")
    calibration_cutoff = int(CALIBRATION_FRACTION * (1 << 64))
    diagnostic_cutoff = int((CALIBRATION_FRACTION + DIAGNOSTIC_FRACTION) * (1 << 64))
    partitions = {"calibration": set(), "diagnostic": set(), "sealed_final": set()}
    for customer_id in sorted(all_ids):
        value = int.from_bytes(
            hashlib.sha256(f"{seed}:{namespace}:{customer_id}".encode()).digest()[:8], "big"
        )
        partition = (
            "calibration"
            if value < calibration_cutoff
            else "diagnostic"
            if value < diagnostic_cutoff
            else "sealed_final"
        )
        partitions[partition].add(customer_id)
    if any(not values for values in partitions.values()):
        raise ValueError("calibration, diagnostic, and sealed-final partitions must be nonempty")
    if set.union(*partitions.values()) != all_ids or any(
        partitions[left] & partitions[right]
        for left in partitions
        for right in partitions
        if left < right
    ):
        raise ValueError("entity-fidelity customer partitions must be exhaustive and disjoint")
    return partitions


def _customer_partition(customer_id: str, *, seed: int = SEED) -> str:
    value = int.from_bytes(
        hashlib.sha256(f"{seed}:uci-real-v2:{customer_id}".encode()).digest()[:8], "big"
    )
    calibration_cutoff = int(CALIBRATION_FRACTION * (1 << 64))
    diagnostic_cutoff = int((CALIBRATION_FRACTION + DIAGNOSTIC_FRACTION) * (1 << 64))
    if value < calibration_cutoff:
        return "calibration"
    return "diagnostic" if value < diagnostic_cutoff else "sealed_final"


def load_retail_partition(
    workbook_path: Path,
    partition: str,
    *,
    seed: int = SEED,
) -> tuple[Any, dict[str, Any]]:
    """Stream one predeclared customer partition; other transaction rows are discarded."""
    if partition not in {"calibration", "diagnostic", "sealed_final"}:
        raise ValueError("unknown entity-fidelity partition")
    provenance = verify_reference_provenance(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    input_rows = 0
    selected_rows: list[dict[str, Any]] = []
    try:
        if workbook.sheetnames != ["Online Retail"]:
            raise ValueError("UCI Online Retail workbook sheet contract changed")
        values = workbook["Online Retail"].iter_rows(values_only=True)
        headers = tuple(next(values))
        if headers != EXPECTED_HEADERS:
            raise ValueError("UCI Online Retail workbook column contract changed")
        customer_index = headers.index("CustomerID")
        for values_row in values:
            input_rows += 1
            customer_id = _customer_id(values_row[customer_index])
            if customer_id is None:
                continue
            assigned = _customer_partition(customer_id, seed=seed)
            if assigned == partition:
                selected_rows.append(dict(zip(headers, values_row, strict=True)))
    finally:
        workbook.close()
    if input_rows != EXPECTED_ROWS:
        raise ValueError("UCI Online Retail workbook line-item count changed")
    frame, counts = aggregate_retail_rows(selected_rows)
    selected_ids_hash = _ids_hash(frame["customer_id"].unique())
    if seed == SEED and selected_ids_hash != EXPECTED_PARTITION_HASHES[partition]:
        raise ValueError("predeclared UCI customer partition membership changed")
    return frame, {
        **provenance,
        **counts,
        "source_input_line_items": input_rows,
        "selected_partition": partition,
        "selected_customer_ids_sha256": selected_ids_hash,
        "nonselected_transactions_materialized": False,
    }


def _empirical_bayes_log_scales(
    groups: list[tuple[str, Any]],
) -> tuple[dict[str, float], float, float]:
    import numpy as np
    from scipy.optimize import minimize
    from scipy.special import gammaln

    observations: list[tuple[str, int, float]] = []
    for customer_id, group in groups:
        values = np.log(group["amount"].to_numpy(dtype=float))
        degrees = len(values) - 1
        residual_sum = float(np.sum((values - values.mean()) ** 2))
        observations.append((str(customer_id), degrees, residual_sum))
    if not observations:
        return {str(customer_id): 0.5 for customer_id, _ in groups}, 2.0, 0.75
    repeats = [(degrees, residual_sum) for _, degrees, residual_sum in observations if degrees]

    def objective(log_parameters: Any) -> float:
        alpha, beta = np.exp(log_parameters)
        return float(
            -sum(
                alpha * math.log(beta)
                - gammaln(alpha)
                + gammaln(alpha + degrees / 2)
                - (alpha + degrees / 2) * math.log(beta + residual_sum / 2)
                for degrees, residual_sum in repeats
            )
        )

    optimum = minimize(
        objective,
        np.log((1.2, 0.25)),
        method="L-BFGS-B",
        bounds=((-10, 10), (-10, 10)),
    )
    if not optimum.success:
        raise RuntimeError("empirical-Bayes log-scale fit did not converge")
    alpha, beta = np.exp(optimum.x)
    estimates = {
        customer_id: math.sqrt(
            (beta + residual_sum / 2) / (alpha + degrees / 2 + 1)
        )
        for customer_id, degrees, residual_sum in observations
    }
    return estimates, float(alpha), float(beta)


def _empirical_bayes_short_probabilities(
    groups: list[tuple[str, Any]],
) -> tuple[dict[str, float], float, float]:
    import numpy as np
    from scipy.optimize import minimize
    from scipy.special import betaln

    stats: list[tuple[str, int, int]] = []
    for customer_id, group in groups:
        gaps = group.sort_values("ts", kind="stable")["ts"].diff().dt.total_seconds().dropna()
        stats.append(
            (
                str(customer_id),
                int(((gaps > 0) & (gaps <= 86_400)).sum()),
                len(gaps),
            )
        )
    successes = np.asarray([row[1] for row in stats], dtype=float)
    trials = np.asarray([row[2] for row in stats], dtype=float)
    global_probability = float(successes.sum() / max(1, trials.sum()))

    def objective(log_parameters: Any) -> float:
        alpha, beta = np.exp(log_parameters)
        return float(
            -np.sum(
                betaln(successes + alpha, trials - successes + beta)
                - betaln(alpha, beta)
            )
        )

    initial_strength = 4.0
    optimum = minimize(
        objective,
        np.log(
            (
                max(1e-3, initial_strength * global_probability),
                max(1e-3, initial_strength * (1 - global_probability)),
            )
        ),
        method="L-BFGS-B",
        bounds=((-10, 10), (-10, 10)),
    )
    if not optimum.success:
        raise RuntimeError("empirical-Bayes short-session fit did not converge")
    alpha, beta = np.exp(optimum.x)
    estimates = {
        customer_id: float((success + alpha) / (trials_count + alpha + beta))
        for customer_id, success, trials_count in stats
    }
    return estimates, float(alpha), float(beta)


def fit_behavior_profile(real_calibration: Any, *, customer_ids_sha256: str) -> BehaviorProfile:
    np, pd, _, _ = _dependencies()
    from scipy.stats import norm, rankdata

    if real_calibration.empty:
        raise ValueError("calibration partition is empty")
    real_calibration = real_calibration.sort_values(["customer_id", "ts"], kind="stable")
    groups = [
        (str(customer_id), group.copy())
        for customer_id, group in real_calibration.groupby("customer_id", sort=True)
    ]
    counts = real_calibration.groupby("customer_id", sort=True).size()
    eb_log_scales, amount_prior_alpha, amount_prior_beta = _empirical_bayes_log_scales(
        groups
    )
    (
        eb_short_probabilities,
        short_prior_alpha,
        short_prior_beta,
    ) = _empirical_bayes_short_probabilities(groups)
    hour_counts = real_calibration["ts"].dt.hour.value_counts()
    hour_weights = tuple(float(hour_counts.get(hour, 0) / len(real_calibration)) for hour in range(24))
    timestamps = real_calibration["ts"]
    window_start = timestamps.min().normalize()
    window_end = timestamps.max().normalize()
    observation_days = max(1.0, float((window_end - window_start).days + 1))
    elapsed_fraction = (timestamps.dt.normalize() - window_start).dt.days / observation_days
    calendar_bins = np.minimum((elapsed_fraction * 12).astype(int), 11)
    calendar_counts = calendar_bins.value_counts()
    calendar_weights = tuple(
        float(calendar_counts.get(index, 0) / len(real_calibration)) for index in range(12)
    )
    all_gaps = _positive_gaps_seconds(real_calibration)
    short_gaps = all_gaps[all_gaps <= 86_400]
    if len(short_gaps) == 0:
        short_gaps = np.asarray(
            [min(86_400.0, float(np.median(all_gaps))) if len(all_gaps) else 3_600.0]
        )
    session_gap_logs = np.log(short_gaps)

    log_amounts = np.log(real_calibration["amount"].to_numpy(dtype=float))
    global_log_center = float(np.mean(log_amounts))
    global_angles = 2 * np.pi * timestamps.dt.hour.to_numpy(dtype=float) / 24
    global_sin = float(np.mean(np.sin(global_angles)))
    global_cos = float(np.mean(np.cos(global_angles)))
    global_resultant = math.hypot(global_sin, global_cos)
    circadian_phase_cut = (
        math.atan2(global_sin, global_cos) + math.pi
    ) % (2 * math.pi)

    primitive_rows: list[list[float]] = []
    singleton_rows: list[list[float]] = []
    repeated_same_day: list[bool] = []
    for customer_id, group in groups:
        group = group.sort_values("ts", kind="stable")
        count = len(group)
        group_logs = np.log(group["amount"].to_numpy(dtype=float))
        angles = 2 * np.pi * group["ts"].dt.hour.to_numpy(dtype=float) / 24
        mean_sin = float(np.mean(np.sin(angles)))
        mean_cos = float(np.mean(np.cos(angles)))
        resultant = math.hypot(mean_sin, mean_cos)
        if resultant > 1e-12:
            phase_sin, phase_cos = mean_sin / resultant, mean_cos / resultant
        else:
            phase_norm = max(1e-12, global_resultant)
            phase_sin, phase_cos = global_sin / phase_norm, global_cos / phase_norm
        phase_latent = (
            math.atan2(phase_sin, phase_cos) - circadian_phase_cut
        ) % (2 * math.pi)
        eb_resultant = min(
            0.999999,
            (count * resultant + EB_CIRCADIAN_PRIOR_STRENGTH * global_resultant)
            / (count + EB_CIRCADIAN_PRIOR_STRENGTH),
        )
        concentration = _resultant_to_concentration(eb_resultant)
        first_day = float((group["ts"].iloc[0].normalize() - window_start).days)
        last_day = float((group["ts"].iloc[-1].normalize() - window_start).days)
        if count == 1:
            singleton_rows.append(
                [
                    first_day / observation_days,
                    float(group_logs.mean() - global_log_center),
                    phase_latent,
                ]
            )
            continue
        same_day = first_day == last_day
        repeated_same_day.append(same_day)
        if not same_day:
            primitive_rows.append(
                [
                    float(count),
                    first_day / observation_days,
                    (last_day - first_day) / observation_days,
                    float(group_logs.mean() - global_log_center),
                    eb_log_scales[customer_id],
                    eb_short_probabilities[customer_id],
                    phase_latent,
                    concentration,
                ]
            )
    primitives = np.asarray(primitive_rows, dtype=float)
    singletons = np.asarray(
        singleton_rows if singleton_rows else [[0.0, 0.0, 0.0]], dtype=float
    )
    normal_scores = np.column_stack(
        [
            norm.ppf((rankdata(primitives[:, index], method="average") - 0.5) / len(primitives))
            for index in range(primitives.shape[1])
        ]
    )
    rank_correlation = np.eye(len(COPULA_PRIMITIVES))
    variable = np.flatnonzero(np.std(normal_scores, axis=0) > 1e-12)
    if len(variable) > 1:
        rank_correlation[np.ix_(variable, variable)] = np.corrcoef(
            normal_scores[:, variable], rowvar=False
        )
    shrunk = (
        (1 - COPULA_CORRELATION_SHRINKAGE) * rank_correlation
        + COPULA_CORRELATION_SHRINKAGE * np.eye(len(COPULA_PRIMITIVES))
    )
    correlation = _nearest_correlation_psd(shrunk)
    quantiles = np.quantile(
        primitives,
        COPULA_QUANTILE_PROBABILITIES,
        axis=0,
        method="linear",
    ).T
    singleton_quantiles = np.quantile(
        singletons,
        COPULA_QUANTILE_PROBABILITIES,
        axis=0,
        method="linear",
    ).T
    return BehaviorProfile(
        source_dataset="UCI Online Retail 352",
        source_partition="calibration",
        source_customer_ids_sha256=customer_ids_sha256,
        source_transactions_sha256=_transactions_hash(real_calibration),
        source_customer_count=len(counts),
        source_transaction_count=len(real_calibration),
        quantile_probabilities=COPULA_QUANTILE_PROBABILITIES,
        primitive_quantiles=tuple(tuple(map(float, row)) for row in quantiles),
        singleton_quantiles=tuple(
            tuple(map(float, row)) for row in singleton_quantiles
        ),
        copula_correlation=tuple(tuple(map(float, row)) for row in correlation),
        hour_weights=hour_weights,
        calendar_weights=calendar_weights,
        observation_days=observation_days,
        session_gap_log_mean=float(np.mean(session_gap_logs)),
        session_gap_log_std=float(np.std(session_gap_logs)),
        repeat_same_day_probability=float(np.mean(repeated_same_day)),
        singleton_probability=len(singleton_rows) / len(groups),
        circadian_phase_cut=circadian_phase_cut,
        amount_variance_prior_alpha=amount_prior_alpha,
        amount_variance_prior_beta=amount_prior_beta,
        short_session_prior_alpha=short_prior_alpha,
        short_session_prior_beta=short_prior_beta,
    )


def _resultant_to_concentration(resultant: float) -> float:
    if resultant < 0.53:
        return 2 * resultant + resultant**3 + 5 * resultant**5 / 6
    if resultant < 0.85:
        return -0.4 + 1.39 * resultant + 0.43 / (1 - resultant)
    return 1 / (resultant**3 - 4 * resultant**2 + 3 * resultant)


def _nearest_correlation_psd(matrix: Any) -> Any:
    np, _, _, _ = _dependencies()
    symmetric = (np.asarray(matrix, dtype=float) + np.asarray(matrix, dtype=float).T) / 2
    eigenvalues, eigenvectors = np.linalg.eigh(symmetric)
    projected = eigenvectors @ np.diag(np.clip(eigenvalues, 1e-10, None)) @ eigenvectors.T
    scale = np.sqrt(np.diag(projected))
    projected = projected / np.outer(scale, scale)
    return (projected + projected.T) / 2


def _positive_gaps_seconds(frame: Any) -> Any:
    np, _, _, _ = _dependencies()
    gaps = frame.groupby("customer_id", sort=False)["ts"].diff().dt.total_seconds().to_numpy()
    return gaps[np.isfinite(gaps) & (gaps > 0)]


@dataclass(frozen=True)
class DomainCalibration:
    amount_log_reference: Any
    window_start: Any
    window_end: Any
    time_scale: float = 1.0


def fit_calibrations(real_calibration: Any, synthetic_calibration: Any) -> tuple[DomainCalibration, DomainCalibration, dict[str, Any]]:
    np, _, _, _ = _dependencies()
    real_gaps = _positive_gaps_seconds(real_calibration)
    synthetic_gaps = _positive_gaps_seconds(synthetic_calibration)
    if len(real_gaps) == 0 or len(synthetic_gaps) == 0:
        raise ValueError("calibration customers need repeated transactions to fit the time scale")
    real_gap_median = float(np.median(real_gaps))
    synthetic_gap_median = float(np.median(synthetic_gaps))
    real_window_seconds = (
        real_calibration["ts"].max() - real_calibration["ts"].min()
    ).total_seconds()
    synthetic_window_seconds = (
        synthetic_calibration["ts"].max() - synthetic_calibration["ts"].min()
    ).total_seconds()
    time_scale = real_window_seconds / synthetic_window_seconds

    def calibration(frame: Any, scale: float) -> DomainCalibration:
        return DomainCalibration(
            amount_log_reference=np.sort(np.log1p(frame["amount"].to_numpy(dtype=float))),
            window_start=frame["ts"].min(),
            window_end=frame["ts"].max(),
            time_scale=scale,
        )

    real = calibration(real_calibration, 1.0)
    synthetic = calibration(synthetic_calibration, time_scale)
    return real, synthetic, {
        "amount_transform": "within-domain ECDF of log1p transaction amount, fitted on calibration customers only",
        "time_transform": "one global multiplier on synthetic elapsed durations, fitted from calibration observation-window duration",
        "real_median_interpurchase_seconds": real_gap_median,
        "synthetic_median_interpurchase_seconds_before_scale": synthetic_gap_median,
        "real_observation_window_seconds": real_window_seconds,
        "synthetic_observation_window_seconds_before_scale": synthetic_window_seconds,
        "synthetic_elapsed_time_scale": time_scale,
    }


def _max_burst(times_seconds: Any, window_seconds: float) -> int:
    left = 0
    best = 0
    for right, timestamp in enumerate(times_seconds):
        while timestamp - times_seconds[left] > window_seconds:
            left += 1
        best = max(best, right - left + 1)
    return best


def customer_features(frame: Any, calibration: DomainCalibration) -> Any:
    np, pd, _, _ = _dependencies()
    rows: list[dict[str, float | str]] = []
    reference = calibration.amount_log_reference
    window_days = max(
        (calibration.window_end - calibration.window_start).total_seconds()
        * calibration.time_scale
        / 86_400,
        1.0,
    )
    for customer_id, group in frame.groupby("customer_id", sort=True):
        group = group.sort_values("ts", kind="stable")
        amounts = group["amount"].to_numpy(dtype=float)
        amount_percentiles = np.searchsorted(reference, np.log1p(amounts), side="right") / len(reference)
        timestamps = group["ts"]
        raw_seconds = timestamps.to_numpy(dtype="datetime64[ns]").astype("int64") / 1e9
        elapsed_seconds = (raw_seconds - raw_seconds[0]) * calibration.time_scale
        gaps = np.diff(elapsed_seconds)
        positive_gaps = gaps[gaps > 0]
        hours = timestamps.dt.hour.to_numpy(dtype=float)
        angles = 2 * np.pi * hours / 24
        recency_days = max(
            0.0,
            (calibration.window_end - timestamps.iloc[-1]).total_seconds()
            * calibration.time_scale
            / 86_400,
        )
        rows.append(
            {
                "customer_id": str(customer_id),
                "median_log_amount_percentile": float(np.median(amount_percentiles)),
                "amount_cv": float(np.std(amounts) / np.mean(amounts)),
                "mean_hour_sin": float(np.mean(np.sin(angles))),
                "mean_hour_cos": float(np.mean(np.cos(angles))),
                "log1p_median_interpurchase_days": float(
                    np.log1p(np.median(positive_gaps) / 86_400) if len(positive_gaps) else np.nan
                ),
                "log1p_frequency_per_30d": float(np.log1p(len(group) * 30 / window_days)),
                "log1p_recency_days": float(np.log1p(recency_days)),
                "log1p_max_burst_24h": float(np.log1p(_max_burst(elapsed_seconds, 86_400))),
            }
        )
    return pd.DataFrame(rows).set_index("customer_id")[list(FEATURES)]


def _ids_hash(ids: Iterable[str]) -> str:
    return hashlib.sha256("\n".join(sorted(ids)).encode()).hexdigest()


def _transactions_hash(frame: Any) -> str:
    digest = hashlib.sha256()
    for row in frame.itertuples(index=False):
        digest.update(f"{row.customer_id}|{row.ts.isoformat()}|{row.amount:.12g}\n".encode())
    return digest.hexdigest()


def _bootstrap_auc(real_predictions: Any, synthetic_predictions: Any, *, seed: int, repeats: int) -> tuple[list[float], float, float]:
    np, _, _, roc_auc_score = _dependencies()
    if repeats < 100:
        raise ValueError("at least 100 stratified customer bootstraps are required")
    rng = np.random.default_rng(seed)
    aucs = []
    labels = np.concatenate((np.zeros(len(real_predictions)), np.ones(len(synthetic_predictions))))
    for _ in range(repeats):
        real_sample = real_predictions[rng.integers(0, len(real_predictions), len(real_predictions))]
        synthetic_sample = synthetic_predictions[
            rng.integers(0, len(synthetic_predictions), len(synthetic_predictions))
        ]
        aucs.append(float(roc_auc_score(labels, np.concatenate((real_sample, synthetic_sample)))))
    low, high = np.quantile(aucs, (0.025, 0.975))
    return aucs, float(low), float(high)


def _classifier(seed: int) -> Any:
    _, _, lightgbm, _ = _dependencies()
    return lightgbm.LGBMClassifier(
        objective="binary",
        n_estimators=150,
        learning_rate=0.04,
        num_leaves=15,
        max_depth=5,
        min_child_samples=40,
        reg_lambda=1.0,
        random_state=seed,
        n_jobs=1,
        deterministic=True,
        force_col_wise=True,
        verbosity=-1,
    )


def _classifier_protocol() -> dict[str, Any]:
    return {
        "implementation": "lightgbm.LGBMClassifier",
        "objective": "binary",
        "n_estimators": 150,
        "learning_rate": 0.04,
        "num_leaves": 15,
        "max_depth": 5,
        "min_child_samples": 40,
        "reg_lambda": 1.0,
        "n_jobs": 1,
        "deterministic": True,
        "force_col_wise": True,
    }


def _canonical_hash(value: Mapping[str, Any]) -> str:
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def _protocol_hash(value: Mapping[str, Any]) -> str:
    return _canonical_hash({key: item for key, item in value.items() if key != "protocol_sha256"})


def _implementation_hashes() -> dict[str, str]:
    return {
        str(path.relative_to(ROOT)): sha256_file(path)
        for path in (ROOT / "sim" / "population.py", ROOT / "sim" / "entity_fidelity.py")
    }


def _write_exclusive_json(path: Path, value: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8") as handle:
        handle.write(json.dumps(value, indent=2, sort_keys=True) + "\n")
        handle.flush()
        os.fsync(handle.fileno())


def freeze_v3_protocol(
    reference_xlsx: Path = DEFAULT_REFERENCE,
    protocol_path: Path = DEFAULT_PROTOCOL,
) -> dict[str, Any]:
    real_calibration, reference = load_retail_partition(reference_xlsx, "calibration")
    calibration_ids = set(real_calibration["customer_id"])
    profile = fit_behavior_profile(
        real_calibration,
        customer_ids_sha256=_ids_hash(calibration_ids),
    )
    protocol: dict[str, Any] = {
        "schema_version": 3,
        "artifact": "entity_fidelity_v3_frozen_protocol",
        "source": {
            "dataset_id": DATASET_ID,
            "source_url": SOURCE_URL,
            "workbook_sha256": reference["workbook_sha256"],
            "manifest_sha256": reference["manifest_sha256"],
            "license": reference["license"],
        },
        "implementation_sha256": _implementation_hashes(),
        "dependency_versions": {
            name: importlib.metadata.version(name)
            for name in ("numpy", "scipy", "pandas", "scikit-learn", "lightgbm", "openpyxl")
        },
        "partition": {
            "assignment": "first 64 bits SHA-256(seed:uci-real-v2:customer_id)",
            "fractions": {
                "calibration": CALIBRATION_FRACTION,
                "diagnostic": DIAGNOSTIC_FRACTION,
                "sealed_final": SEALED_FINAL_FRACTION,
            },
            "customer_ids_sha256": EXPECTED_PARTITION_HASHES,
            "diagnostic_permitted_for_v3": False,
        },
        "generator": {
            "primitives": list(COPULA_PRIMITIVES),
            "singleton_primitives": list(SINGLETON_PRIMITIVES),
            "marginal_quantile_probabilities": list(COPULA_QUANTILE_PROBABILITIES),
            "rank_transform": "average ranks, u=(rank-0.5)/n, standard-normal inverse CDF",
            "correlation_shrinkage_toward_identity": COPULA_CORRELATION_SHRINKAGE,
            "psd_projection": "symmetric eigendecomposition; eigenvalues clipped to 1e-10; unit-diagonal rescale",
            "empirical_bayes": {
                "within_customer_log_amount_scale": "inverse-gamma variance prior alpha/beta by calibration marginal MLE; per-customer posterior mode sqrt((beta+SSE/2)/(alpha+(n-1)/2+1))",
                "short_session_probability": "beta-binomial EB; alpha/beta fitted by calibration marginal likelihood",
                "circadian_resultant_prior_strength": EB_CIRCADIAN_PRIOR_STRENGTH,
            },
            "mixed_support": "calibration-rate singleton hurdle; singleton count=1/duration=0/discrete supported-hour phase; repeat-only copula with calibrated Bernoulli same-day atom",
            "schedule": "exact largest-remainder customer quotas; persistent active-window geometric sessions; calibration calendar dates; phase/concentration tilted global hours; <=24h calibration lognormal session gaps; chronological heap merge",
            "contains_customer_ids_vectors_raw_rows_or_evaluator_features": False,
            "behavior_profile": profile.as_dict(),
            "behavior_profile_sha256": profile.sha256,
        },
        "evaluation": {
            "seeds": list(READINESS_SEEDS),
            "cross_fit_folds": CROSS_FIT_FOLDS,
            "cross_fit_assignment": "first 64 bits SHA-256(seed:uci-v3-crossfit:customer_id) modulo 5",
            "features": list(FEATURES),
            "domain_transforms": {
                "amount": "within-domain calibration-only ECDF of log1p transaction amount",
                "time": "calibration-only synthetic-to-real observation-window duration ratio",
            },
            "classifier": _classifier_protocol(),
            "orientation_free_statistic": "D=max(AUC,1-AUC)",
            "readiness": "max across seeds of fold-customer-count-weighted D",
            "bootstrap": "10,000 stratified customer bootstraps; 2.5/97.5 percentiles",
            "readiness_bootstrap_repeats": READINESS_BOOTSTRAP_REPEATS,
            "final_bootstrap": "shared real bootstrap indices across seeds; independent synthetic indices",
            "threshold_max_upper_95": GATE_MAX_AUC_UPPER_95,
        },
        "sealed_final": {
            "confirmation": f"{SEALED_FINAL_CONFIRMATION}:<protocol_sha256>",
            "exclusive_receipt_before_partition_selection": True,
            "second_invocation_fails": True,
        },
    }
    protocol["protocol_sha256"] = _protocol_hash(protocol)
    if protocol_path.exists():
        existing = json.loads(protocol_path.read_text(encoding="utf-8"))
        if existing != protocol:
            raise FileExistsError("frozen v3 protocol exists with different content")
        return existing
    _write_exclusive_json(protocol_path, protocol)
    protocol_path.chmod(0o444)
    return protocol


def load_frozen_protocol(protocol_path: Path = DEFAULT_PROTOCOL) -> dict[str, Any]:
    protocol = json.loads(protocol_path.read_text(encoding="utf-8"))
    if protocol.get("protocol_sha256") != _protocol_hash(protocol):
        raise ValueError("frozen v3 protocol hash mismatch")
    if protocol.get("implementation_sha256") != _implementation_hashes():
        raise ValueError("implementation changed after the v3 protocol was frozen")
    if protocol["source"]["workbook_sha256"] != sha256_file(DEFAULT_REFERENCE):
        raise ValueError("reference workbook changed after the v3 protocol was frozen")
    return protocol


def _synthetic_group(
    profile: BehaviorProfile,
    *,
    customers: int,
    seed: int,
    prefix: str,
) -> Any:
    frame = generate_synthetic_transactions(
        events=None,
        customers=customers,
        seed=seed,
        behavior_profile=profile,
    )
    frame["customer_id"] = prefix + frame["customer_id"].astype(str)
    return frame


def _crossfit_seed(
    real_calibration: Any,
    *,
    seed: int,
    fold_numbers: Iterable[int] | None = None,
) -> list[dict[str, Any]]:
    np, pd, _, roc_auc_score = _dependencies()
    customer_ids = sorted(set(real_calibration["customer_id"]))
    fold_by_id = {
        customer_id: int.from_bytes(
            hashlib.sha256(f"{seed}:uci-v3-crossfit:{customer_id}".encode()).digest()[:8],
            "big",
        )
        % CROSS_FIT_FOLDS
        for customer_id in customer_ids
    }
    folds: list[dict[str, Any]] = []
    for fold in fold_numbers if fold_numbers is not None else range(CROSS_FIT_FOLDS):
        evaluation_ids = {customer_id for customer_id, value in fold_by_id.items() if value == fold}
        training_ids = set(customer_ids) - evaluation_ids
        if not evaluation_ids or not training_ids:
            raise ValueError("cross-fit fold is empty")
        real_training = real_calibration.loc[
            real_calibration["customer_id"].isin(training_ids)
        ].copy()
        real_evaluation = real_calibration.loc[
            real_calibration["customer_id"].isin(evaluation_ids)
        ].copy()
        profile = fit_behavior_profile(
            real_training,
            customer_ids_sha256=_ids_hash(training_ids),
        )
        synthetic_training = _synthetic_group(
            profile,
            customers=len(training_ids),
            seed=seed + 10_000 * fold + 101,
            prefix=f"train-{seed}-{fold}-",
        )
        synthetic_evaluation = _synthetic_group(
            profile,
            customers=len(evaluation_ids),
            seed=seed + 10_000 * fold + 202,
            prefix=f"eval-{seed}-{fold}-",
        )
        real_transform, synthetic_transform, _ = fit_calibrations(
            real_training, synthetic_training
        )
        real_training_features = customer_features(real_training, real_transform)
        synthetic_training_features = customer_features(
            synthetic_training, synthetic_transform
        )
        x_train = pd.concat(
            (real_training_features, synthetic_training_features), ignore_index=True
        )
        y_train = np.concatenate(
            (
                np.zeros(len(real_training_features), dtype="int8"),
                np.ones(len(synthetic_training_features), dtype="int8"),
            )
        )
        classifier = _classifier(seed + fold)
        classifier.fit(x_train, y_train)
        real_predictions = classifier.predict_proba(
            customer_features(real_evaluation, real_transform)
        )[:, 1]
        synthetic_predictions = classifier.predict_proba(
            customer_features(synthetic_evaluation, synthetic_transform)
        )[:, 1]
        labels = np.concatenate(
            (np.zeros(len(real_predictions)), np.ones(len(synthetic_predictions)))
        )
        auc = float(
            roc_auc_score(labels, np.concatenate((real_predictions, synthetic_predictions)))
        )
        folds.append(
            {
                "fold": fold,
                "training_real_customers": len(training_ids),
                "evaluation_real_customers": len(evaluation_ids),
                "evaluation_synthetic_customers": len(synthetic_predictions),
                "training_customer_ids_sha256": _ids_hash(training_ids),
                "evaluation_customer_ids_sha256": _ids_hash(evaluation_ids),
                "profile_sha256": profile.sha256,
                "auc": auc,
                "d": max(auc, 1 - auc),
                "_real_predictions": real_predictions,
                "_synthetic_predictions": synthetic_predictions,
            }
        )
    return folds


def _orientation_free_bootstrap(
    seed_folds: Mapping[int, list[dict[str, Any]]],
    *,
    repeats: int,
    seed: int,
) -> tuple[list[float], float, float]:
    np, _, _, roc_auc_score = _dependencies()
    if repeats < 100:
        raise ValueError("at least 100 stratified customer bootstraps are required")
    rng = np.random.default_rng(seed)
    maxima: list[float] = []
    for _ in range(repeats):
        seed_values = []
        for folds in seed_folds.values():
            weighted = 0.0
            total_weight = 0
            for fold in folds:
                real = fold["_real_predictions"]
                synthetic = fold["_synthetic_predictions"]
                sampled_real = real[rng.integers(0, len(real), len(real))]
                sampled_synthetic = synthetic[
                    rng.integers(0, len(synthetic), len(synthetic))
                ]
                labels = np.concatenate((np.zeros(len(real)), np.ones(len(synthetic))))
                auc = float(
                    roc_auc_score(
                        labels, np.concatenate((sampled_real, sampled_synthetic))
                    )
                )
                weight = fold["evaluation_real_customers"]
                weighted += weight * max(auc, 1 - auc)
                total_weight += weight
            seed_values.append(weighted / total_weight)
        maxima.append(max(seed_values))
    low, high = np.quantile(maxima, (0.025, 0.975))
    return maxima, float(low), float(high)


def run_calibration_readiness(
    reference_xlsx: Path = DEFAULT_REFERENCE,
    output: Path = DEFAULT_READINESS,
    *,
    protocol_path: Path = DEFAULT_PROTOCOL,
) -> dict[str, Any]:
    np, _, _, _ = _dependencies()
    protocol = load_frozen_protocol(protocol_path)
    real_calibration, reference = load_retail_partition(reference_xlsx, "calibration")
    profile = fit_behavior_profile(
        real_calibration,
        customer_ids_sha256=_ids_hash(real_calibration["customer_id"].unique()),
    )
    if profile.sha256 != protocol["generator"]["behavior_profile_sha256"]:
        raise ValueError("all-calibration behavior profile differs from frozen protocol")
    seed_folds = {
        seed: _crossfit_seed(real_calibration, seed=seed) for seed in READINESS_SEEDS
    }
    seed_point = {
        str(seed): sum(
            fold["d"] * fold["evaluation_real_customers"] for fold in folds
        )
        / sum(fold["evaluation_real_customers"] for fold in folds)
        for seed, folds in seed_folds.items()
    }
    bootstrap, ci_low, ci_high = _orientation_free_bootstrap(
        seed_folds,
        repeats=READINESS_BOOTSTRAP_REPEATS,
        seed=SEED + 30,
    )
    result = {
        "schema_version": 3,
        "artifact": "entity_fidelity_v3_calibration_crossfit_readiness",
        "protocol_sha256": protocol["protocol_sha256"],
        "calibration_only": True,
        "diagnostic_transactions_accessed": False,
        "sealed_final_transactions_accessed": False,
        "reference": {
            "workbook_sha256": reference["workbook_sha256"],
            "partition": "calibration",
            "customer_ids_sha256": EXPECTED_PARTITION_HASHES["calibration"],
            "customers": int(real_calibration["customer_id"].nunique()),
            "transactions": len(real_calibration),
        },
        "seeds": {
            str(seed): {
                "fold_weighted_d": seed_point[str(seed)],
                "folds": [
                    {key: value for key, value in fold.items() if not key.startswith("_")}
                    for fold in folds
                ],
            }
            for seed, folds in seed_folds.items()
        },
        "metrics": {
            "orientation_free_max_seed_d": max(seed_point.values()),
            "bootstrap_max_seed_d_mean": float(np.mean(bootstrap)),
            "bootstrap_max_seed_d_ci_95": [ci_low, ci_high],
            "threshold_max_d_upper_95": GATE_MAX_AUC_UPPER_95,
        },
        "gate_passed": ci_high <= GATE_MAX_AUC_UPPER_95,
    }
    _write_exclusive_json(output, result)
    return result


def _final_seed_predictions(
    real_calibration: Any,
    real_sealed: Any,
    profile: BehaviorProfile,
    *,
    seed: int,
) -> dict[str, Any]:
    np, pd, _, roc_auc_score = _dependencies()
    training_customers = int(real_calibration["customer_id"].nunique())
    evaluation_customers = int(real_sealed["customer_id"].nunique())
    synthetic_training = _synthetic_group(
        profile,
        customers=training_customers,
        seed=seed + 700_001,
        prefix=f"final-train-{seed}-",
    )
    synthetic_evaluation = _synthetic_group(
        profile,
        customers=evaluation_customers,
        seed=seed + 700_002,
        prefix=f"final-eval-{seed}-",
    )
    real_transform, synthetic_transform, _ = fit_calibrations(
        real_calibration, synthetic_training
    )
    real_training_features = customer_features(real_calibration, real_transform)
    synthetic_training_features = customer_features(
        synthetic_training, synthetic_transform
    )
    classifier = _classifier(seed)
    classifier.fit(
        pd.concat(
            (real_training_features, synthetic_training_features), ignore_index=True
        ),
        np.concatenate(
            (
                np.zeros(len(real_training_features), dtype="int8"),
                np.ones(len(synthetic_training_features), dtype="int8"),
            )
        ),
    )
    real_predictions = classifier.predict_proba(
        customer_features(real_sealed, real_transform)
    )[:, 1]
    synthetic_predictions = classifier.predict_proba(
        customer_features(synthetic_evaluation, synthetic_transform)
    )[:, 1]
    labels = np.concatenate(
        (np.zeros(len(real_predictions)), np.ones(len(synthetic_predictions)))
    )
    auc = float(
        roc_auc_score(labels, np.concatenate((real_predictions, synthetic_predictions)))
    )
    return {
        "auc": auc,
        "d": max(auc, 1 - auc),
        "real_customers": len(real_predictions),
        "synthetic_customers": len(synthetic_predictions),
        "_real_predictions": real_predictions,
        "_synthetic_predictions": synthetic_predictions,
    }


def _final_bootstrap(
    seed_predictions: Mapping[int, dict[str, Any]],
    *,
    repeats: int,
    seed: int,
) -> tuple[list[float], float, float]:
    np, _, _, roc_auc_score = _dependencies()
    rng = np.random.default_rng(seed)
    first = next(iter(seed_predictions.values()))
    real_count = len(first["_real_predictions"])
    if any(len(value["_real_predictions"]) != real_count for value in seed_predictions.values()):
        raise ValueError("final real prediction vectors must share one sealed population")
    maxima = []
    for _ in range(repeats):
        shared_real_indices = rng.integers(0, real_count, real_count)
        seed_values = []
        for value in seed_predictions.values():
            real = value["_real_predictions"][shared_real_indices]
            synthetic_source = value["_synthetic_predictions"]
            synthetic = synthetic_source[
                rng.integers(0, len(synthetic_source), len(synthetic_source))
            ]
            labels = np.concatenate((np.zeros(len(real)), np.ones(len(synthetic))))
            auc = float(roc_auc_score(labels, np.concatenate((real, synthetic))))
            seed_values.append(max(auc, 1 - auc))
        maxima.append(max(seed_values))
    low, high = np.quantile(maxima, (0.025, 0.975))
    return maxima, float(low), float(high)


def run_entity_fidelity_v3_final(
    reference_xlsx: Path = DEFAULT_REFERENCE,
    output: Path = DEFAULT_OUTPUT,
    *,
    protocol_path: Path = DEFAULT_PROTOCOL,
    readiness_path: Path = DEFAULT_READINESS,
    receipt_path: Path = DEFAULT_SEAL_RECEIPT,
    sealed_final_confirmation: str | None = None,
) -> dict[str, Any]:
    np, _, _, _ = _dependencies()
    protocol = load_frozen_protocol(protocol_path)
    expected_confirmation = f"{SEALED_FINAL_CONFIRMATION}:{protocol['protocol_sha256']}"
    if sealed_final_confirmation != expected_confirmation:
        raise ValueError("sealed-final scoring requires the protocol-hash-bound one-time approval")
    readiness = json.loads(readiness_path.read_text(encoding="utf-8"))
    if (
        readiness.get("protocol_sha256") != protocol["protocol_sha256"]
        or readiness.get("gate_passed") is not True
        or readiness.get("calibration_only") is not True
        or readiness.get("diagnostic_transactions_accessed") is not False
        or readiness.get("sealed_final_transactions_accessed") is not False
    ):
        raise ValueError("calibration-only readiness is absent, failed, or not protocol-bound")
    if output.exists():
        raise FileExistsError("final entity-fidelity output already exists")
    if receipt_path.exists():
        raise FileExistsError("sealed-final one-shot receipt already exists")

    real_calibration, _ = load_retail_partition(reference_xlsx, "calibration")
    profile = fit_behavior_profile(
        real_calibration,
        customer_ids_sha256=_ids_hash(real_calibration["customer_id"].unique()),
    )
    if profile.sha256 != protocol["generator"]["behavior_profile_sha256"]:
        raise ValueError("all-calibration profile differs from frozen protocol")
    receipt = {
        "schema_version": 3,
        "artifact": "entity_fidelity_v3_sealed_final_one_shot_receipt",
        "protocol_sha256": protocol["protocol_sha256"],
        "status": "consumed",
        "consumed_at_utc": datetime.now(timezone.utc).isoformat(),
        "confirmation_sha256": hashlib.sha256(
            sealed_final_confirmation.encode()
        ).hexdigest(),
        "sealed_partition_customer_ids_sha256": EXPECTED_PARTITION_HASHES["sealed_final"],
        "readiness_artifact_sha256": sha256_file(readiness_path),
    }
    receipt["receipt_sha256"] = _canonical_hash(receipt)
    _write_exclusive_json(receipt_path, receipt)

    real_sealed, sealed_reference = load_retail_partition(reference_xlsx, "sealed_final")
    seed_predictions = {
        seed: _final_seed_predictions(
            real_calibration, real_sealed, profile, seed=seed
        )
        for seed in READINESS_SEEDS
    }
    bootstrap, ci_low, ci_high = _final_bootstrap(
        seed_predictions,
        repeats=READINESS_BOOTSTRAP_REPEATS,
        seed=SEED + 31,
    )
    point_d = max(value["d"] for value in seed_predictions.values())
    result = {
        "schema_version": 3,
        "artifact": "entity_fidelity_v3_sealed_final",
        "protocol_sha256": protocol["protocol_sha256"],
        "scope": "Entity/behavioral fidelity of legitimate purchase timing and value only; not fraud, authorization, merchant, graph, text, or intent fidelity.",
        "limitations": [
            "UCI Online Retail is one UK non-store retailer from 2010-2011, not a payment-network population.",
            "Filtered invoices are retail purchases, while canonical synthetic events represent payment events.",
            "Amounts use within-domain calibration transforms because GBP and INR levels are not semantically comparable.",
            "UCI InvoiceDate has no published timezone; hour-of-day is used as recorded.",
            "No fraud label exists in this gate and no claim about fraud realism is made.",
        ],
        "protocol": {
            "protocol_sha256": protocol["protocol_sha256"],
            "evaluation_partition": "sealed_final",
            "sealed_final_scored": True,
            "readiness_gate_passed": True,
            "orientation_free_statistic": "D=max(AUC,1-AUC)",
            "bootstrap_repeats": READINESS_BOOTSTRAP_REPEATS,
            "shared_real_bootstrap_indices_across_seeds": True,
            "independent_synthetic_bootstrap_indices_across_seeds": True,
        },
        "bindings": {
            "readiness_artifact_sha256": sha256_file(readiness_path),
            "sealed_receipt_sha256": receipt["receipt_sha256"],
            "behavior_profile_sha256": profile.sha256,
        },
        "behavior_profile": profile.as_dict(),
        "reference": {
            "workbook_sha256": sealed_reference["workbook_sha256"],
            "calibration_customer_ids_sha256": EXPECTED_PARTITION_HASHES["calibration"],
            "sealed_final_customer_ids_sha256": EXPECTED_PARTITION_HASHES["sealed_final"],
            "calibration_customers": int(real_calibration["customer_id"].nunique()),
            "sealed_final_customers": int(real_sealed["customer_id"].nunique()),
            "diagnostic_transactions_accessed": False,
        },
        "seeds": {
            str(seed): {
                key: item
                for key, item in value.items()
                if not key.startswith("_")
            }
            for seed, value in seed_predictions.items()
        },
        "metrics": {
            "orientation_free_max_seed_d": point_d,
            "bootstrap_max_seed_d_mean": float(np.mean(bootstrap)),
            "bootstrap_max_seed_d_ci_95": [ci_low, ci_high],
            "threshold_max_d_upper_95": GATE_MAX_AUC_UPPER_95,
            "gate_passed": ci_high <= GATE_MAX_AUC_UPPER_95,
        },
    }
    _write_exclusive_json(output, result)
    return result


def verify_v3_artifacts(
    *,
    protocol_path: Path = DEFAULT_PROTOCOL,
    readiness_path: Path = DEFAULT_READINESS,
    receipt_path: Path = DEFAULT_SEAL_RECEIPT,
    final_path: Path = DEFAULT_OUTPUT,
) -> dict[str, Any]:
    protocol = load_frozen_protocol(protocol_path)
    readiness = json.loads(readiness_path.read_text(encoding="utf-8"))
    checks = {
        "readiness_protocol_bound": readiness.get("protocol_sha256")
        == protocol["protocol_sha256"],
        "readiness_gate_passed": readiness.get("gate_passed") is True,
        "readiness_calibration_only": readiness.get("calibration_only") is True
        and readiness.get("diagnostic_transactions_accessed") is False
        and readiness.get("sealed_final_transactions_accessed") is False,
        "sealed_artifacts_complete": receipt_path.exists() and final_path.exists(),
    }
    if receipt_path.exists() or final_path.exists():
        if not receipt_path.exists() or not final_path.exists():
            pass
        else:
            receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
            final = json.loads(final_path.read_text(encoding="utf-8"))
            checks.update(
                {
                    "receipt_protocol_bound": receipt.get("protocol_sha256")
                    == protocol["protocol_sha256"],
                    "receipt_consumed": receipt.get("status") == "consumed",
                    "receipt_hash_valid": receipt.get("receipt_sha256")
                    == _canonical_hash(
                        {
                            key: value
                            for key, value in receipt.items()
                            if key != "receipt_sha256"
                        }
                    ),
                    "final_protocol_bound": final.get("protocol", {}).get(
                        "protocol_sha256"
                    )
                    == protocol["protocol_sha256"],
                    "final_is_sealed": final.get("protocol", {}).get(
                        "evaluation_partition"
                    )
                    == "sealed_final"
                    and final.get("protocol", {}).get("sealed_final_scored") is True,
                    "final_readiness_bound": final.get("protocol", {}).get(
                        "readiness_gate_passed"
                    )
                    is True,
                    "final_gate_passed": final.get("metrics", {}).get("gate_passed")
                    is True,
                }
            )
    return {
        "protocol_sha256": protocol["protocol_sha256"],
        "checks": checks,
        "verified": all(checks.values()),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    freeze = commands.add_parser("freeze", help="freeze the calibration-only v3 protocol")
    freeze.add_argument("--reference-xlsx", type=Path, default=DEFAULT_REFERENCE)
    freeze.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    readiness = commands.add_parser(
        "readiness", help="run the calibration-only 5-fold/3-seed readiness gate"
    )
    readiness.add_argument("--reference-xlsx", type=Path, default=DEFAULT_REFERENCE)
    readiness.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    readiness.add_argument("--output", type=Path, default=DEFAULT_READINESS)
    sealed = commands.add_parser(
        "sealed", help="consume the one-shot receipt and score sealed_final exactly once"
    )
    sealed.add_argument("--reference-xlsx", type=Path, default=DEFAULT_REFERENCE)
    sealed.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    sealed.add_argument("--readiness", type=Path, default=DEFAULT_READINESS)
    sealed.add_argument("--receipt", type=Path, default=DEFAULT_SEAL_RECEIPT)
    sealed.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    sealed.add_argument("--confirmation", required=True)
    verify = commands.add_parser("verify", help="verify protocol and artifact bindings")
    verify.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    verify.add_argument("--readiness", type=Path, default=DEFAULT_READINESS)
    verify.add_argument("--receipt", type=Path, default=DEFAULT_SEAL_RECEIPT)
    verify.add_argument("--final", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if args.command == "freeze":
        result = freeze_v3_protocol(args.reference_xlsx, args.protocol)
    elif args.command == "readiness":
        result = run_calibration_readiness(
            args.reference_xlsx, args.output, protocol_path=args.protocol
        )
    elif args.command == "sealed":
        result = run_entity_fidelity_v3_final(
            args.reference_xlsx,
            args.output,
            protocol_path=args.protocol,
            readiness_path=args.readiness,
            receipt_path=args.receipt,
            sealed_final_confirmation=args.confirmation,
        )
    else:
        result = verify_v3_artifacts(
            protocol_path=args.protocol,
            readiness_path=args.readiness,
            receipt_path=args.receipt,
            final_path=args.final,
        )
    print(json.dumps(result, sort_keys=True))
    if args.command == "readiness" and result["gate_passed"] is not True:
        raise SystemExit(1)
    if args.command == "sealed" and result["metrics"]["gate_passed"] is not True:
        raise SystemExit(1)
    if args.command == "verify" and result["verified"] is not True:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
