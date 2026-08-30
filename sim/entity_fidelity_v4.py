"""Fail-closed schema-4 entity-fidelity protocol and evaluation gates."""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import math
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook

from scripts.download_retail_ii_reference import (
    ARCHIVE_SHA256 as UCI502_ARCHIVE_SHA256,
    DATASET_ID as UCI502_DATASET_ID,
    LICENSE as UCI502_LICENSE,
    READINESS_SLICE,
    SOURCE_URL as UCI502_SOURCE_URL,
    WORKBOOK_SHA256 as UCI502_WORKBOOK_SHA256,
    verify_existing as verify_uci502_existing,
)
from scripts.download_retail_reference import (
    DATASET_ID as UCI352_DATASET_ID,
    LICENSE as UCI352_LICENSE,
    SOURCE_URL as UCI352_SOURCE_URL,
    WORKBOOK_SHA256 as UCI352_WORKBOOK_SHA256,
    sha256_file,
)
from sim.entity_fidelity import (
    EB_CIRCADIAN_PRIOR_STRENGTH,
    EXPECTED_PARTITION_HASHES,
    FEATURES,
    DomainCalibration,
    _classifier,
    _classifier_protocol,
    _customer_id,
    _dependencies,
    _empirical_bayes_log_scales,
    _empirical_bayes_short_probabilities,
    _final_bootstrap,
    _ids_hash,
    _positive_gaps_seconds,
    _resultant_to_concentration,
    _transactions_hash,
    _write_exclusive_json,
    aggregate_retail_rows,
    customer_features,
    fit_calibrations,
    load_retail_partition,
)
from sim.population_v4 import (
    CHECKERBOARD_BINS,
    QUANTILE_PROBABILITIES,
    STRATUM_CELL_COUNTS,
    STRATUM_ORDER,
    STRATUM_PRIMITIVES,
    BehaviorProfileV4,
    CheckerboardStratum,
    PopulationSimulatorV4,
)


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_REFERENCE = ROOT / "data" / "reference" / "online_retail.xlsx"
DEFAULT_READINESS_REFERENCE = ROOT / "data" / "reference" / "online_retail_ii.xlsx"
DEFAULT_PROTOCOL = ROOT / "results" / "entity_fidelity_v4_protocol.json"
DEFAULT_EXPLORATORY = ROOT / "results" / "entity_fidelity_v4_exploratory.json"
DEFAULT_READINESS = ROOT / "results" / "entity_fidelity_v4_readiness.json"
DEFAULT_SEAL_RECEIPT = ROOT / "results" / "entity_fidelity_v4_sealed_receipt.json"
DEFAULT_OUTPUT = ROOT / "results" / "entity_fidelity.json"
SEED = 20260827
EVALUATION_SEEDS = (20260827, 20260828, 20260829)
CROSS_FIT_FOLDS = 5
BOOTSTRAP_REPEATS = 10_000
GATE_MAX_D_UPPER_95 = 0.65
NOVELTY_DRAWS = 100_000
MIN_STRATUM_CUSTOMERS = 32
SEALED_FINAL_CONFIRMATION = "APPROVED_SINGLE_FINAL_SCORE"
UCI502_SHEET = str(READINESS_SLICE["sheet"])
UCI502_EXPECTED_ROWS = 525_461
UCI502_HEADERS = (
    "Invoice",
    "StockCode",
    "Description",
    "Quantity",
    "InvoiceDate",
    "Price",
    "Customer ID",
    "Country",
)
UCI502_WINDOW_START = "2009-12-01"
UCI502_WINDOW_END = str(READINESS_SLICE["invoice_date_before"])
UCI502_SPLIT_NAMESPACE = "20260827:uci502-v4-readiness-v1"
UCI502_MIN_SPLIT_CUSTOMERS = 500
UCI502_SPLIT_ASSIGNMENT = (
    "first 64 SHA-256 bits of 20260827:uci502-v4-readiness-v1:<customer_id>; "
    "<2^63 transform-only, otherwise score"
)
EXPLORATORY_ASSIGNMENT = (
    "SHA256(seed:uci-v4-crossfit:customer_id) first64 modulo 5"
)
CHECKERBOARD_TIE_RULE = (
    "I=[count(x'<x)/n,count(x'<=x)/n]; cell mass is the mean product of "
    "fractional overlaps with thirds"
)
CHECKERBOARD_CELL_ALLOCATION = (
    "floor(m*p) plus residual resampling from normalized fractional residuals"
)
CHECKERBOARD_SAMPLING = (
    "U=(bin+V)/3 with fresh independent V; linear inverse through frozen "
    "32-knot marginals"
)
FORBIDDEN_ARTIFACT_KEYS = {
    "customer_ids",
    "customer_vectors",
    "individual_hashes",
    "raw_rows",
    "transaction_rows",
    "timestamps",
    "rank_rows",
    "rank_matrix",
    "feature_rows",
    "evaluator_feature_rows",
    "predictions",
    "prediction_vectors",
}
AGGREGATE_CONTAINER_ALLOWLIST = {
    "implementation_sha256",
    "dependency_versions",
    "sources",
    "training_and_sealed",
    "temporal_readiness",
    "generator",
    "strata",
    "novelty",
    "evaluation",
    "features",
    "classifier",
    "seeds",
    "sealed_final",
    "source",
    "checkerboard",
    "quantile_probabilities",
    "primitive_quantiles",
    "primitives",
    "cell_masses",
    "renderer",
    "hour_weights",
    "calendar_weights",
    "empirical_bayes",
    "privacy",
    "reference",
    "folds",
    "metrics",
    "bootstrap_max_seed_d_ci_95",
    "protocol",
    "bindings",
    "behavior_profile",
}
CHECKERBOARD_PRIMITIVE_NAMES = frozenset(
    primitive
    for primitives in STRATUM_PRIMITIVES.values()
    for primitive in primitives
)


def _canonical_hash(value: Mapping[str, Any]) -> str:
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, separators=(",", ":")).encode()
    ).hexdigest()


def _protocol_hash(value: Mapping[str, Any]) -> str:
    return _canonical_hash(
        {key: item for key, item in value.items() if key != "protocol_sha256"}
    )


def _implementation_hashes() -> dict[str, str]:
    paths = (
        ROOT / "sim" / "population.py",
        ROOT / "sim" / "entity_fidelity.py",
        ROOT / "sim" / "population_v4.py",
        ROOT / "sim" / "entity_fidelity_v4.py",
        ROOT / "scripts" / "download_retail_reference.py",
        ROOT / "scripts" / "download_retail_ii_reference.py",
    )
    return {str(path.relative_to(ROOT)): sha256_file(path) for path in paths}


def _assert_aggregate_only(value: Any, *, path: str = "artifact") -> None:
    """Allow only predeclared aggregate containers in persisted v4 artifacts."""
    if isinstance(value, Mapping):
        for key, item in value.items():
            normalized = str(key).lower()
            if normalized in FORBIDDEN_ARTIFACT_KEYS:
                raise ValueError(f"{path}.{key} is not allowed in a v4 artifact")
            if isinstance(item, (Mapping, list, tuple)) and not (
                normalized in AGGREGATE_CONTAINER_ALLOWLIST
                or normalized in STRATUM_ORDER
                or normalized in CHECKERBOARD_PRIMITIVE_NAMES
                or normalized.isdigit()
            ):
                raise ValueError(
                    f"{path}.{key} is not an allowlisted aggregate container"
                )
            _assert_aggregate_only(item, path=f"{path}.{key}")
    elif isinstance(value, (list, tuple)):
        for index, item in enumerate(value):
            _assert_aggregate_only(item, path=f"{path}[{index}]")


def _write_artifact(path: Path, value: Mapping[str, Any]) -> None:
    _assert_aggregate_only(value)
    _write_exclusive_json(path, value)
    path.chmod(0o444)


def _gate_metrics_valid(value: Mapping[str, Any]) -> bool:
    metrics = value.get("metrics", {})
    interval = metrics.get("bootstrap_max_seed_d_ci_95", ())
    return (
        value.get("gate_passed") is True
        and metrics.get("threshold_max_d_upper_95") == GATE_MAX_D_UPPER_95
        and isinstance(interval, list)
        and len(interval) == 2
        and all(isinstance(item, (int, float)) and math.isfinite(item) for item in interval)
        and interval[0] <= interval[1] <= GATE_MAX_D_UPPER_95
    )


def fractional_tie_weights(values: Any, *, bins: int = CHECKERBOARD_BINS) -> Any:
    """Fractional CDF tie-interval overlap with equal-width rank bins."""
    np, _, _, _ = _dependencies()
    observations = np.asarray(values, dtype=float)
    if observations.ndim != 1 or len(observations) == 0 or not np.isfinite(observations).all():
        raise ValueError("fractional tie weights need a finite nonempty vector")
    ordered = np.sort(observations)
    lower = np.searchsorted(ordered, observations, side="left") / len(observations)
    upper = np.searchsorted(ordered, observations, side="right") / len(observations)
    width = upper - lower
    weights = np.zeros((len(observations), bins), dtype=float)
    for bin_index in range(bins):
        bin_lower = bin_index / bins
        bin_upper = (bin_index + 1) / bins
        overlap = np.maximum(
            0.0, np.minimum(upper, bin_upper) - np.maximum(lower, bin_lower)
        )
        weights[:, bin_index] = overlap / width
    if not np.allclose(weights.sum(axis=1), 1.0, atol=1e-12):
        raise RuntimeError("fractional tie weights did not partition every interval")
    return weights


def fractional_checkerboard_masses(
    primitive_rows: Any, *, bins: int = CHECKERBOARD_BINS
) -> tuple[float, ...]:
    """Compute fixed-size checkerboard masses without smoothing or stored ranks."""
    np, _, _, _ = _dependencies()
    rows = np.asarray(primitive_rows, dtype=float)
    if rows.ndim != 2 or rows.shape[0] == 0 or rows.shape[1] == 0:
        raise ValueError("checkerboard fit needs a nonempty primitive matrix")
    coordinate_weights = [
        fractional_tie_weights(rows[:, index], bins=bins)
        for index in range(rows.shape[1])
    ]
    masses = np.zeros(bins ** rows.shape[1], dtype=float)
    for row_index in range(rows.shape[0]):
        contribution = coordinate_weights[0][row_index]
        for coordinate in range(1, rows.shape[1]):
            contribution = np.multiply.outer(
                contribution, coordinate_weights[coordinate][row_index]
            )
        masses += contribution.reshape(-1)
    masses /= rows.shape[0]
    if not math.isclose(float(masses.sum()), 1.0, abs_tol=1e-12):
        raise RuntimeError("checkerboard masses do not sum to one")
    return tuple(map(float, masses))


def _fit_behavior_profile_and_rows(
    real_calibration: Any, *, customer_ids_sha256: str
) -> tuple[BehaviorProfileV4, dict[str, Any]]:
    np, _, _, _ = _dependencies()
    if real_calibration.empty:
        raise ValueError("calibration partition is empty")
    frame = real_calibration.sort_values(["customer_id", "ts"], kind="stable")
    groups = [
        (str(customer_id), group.copy())
        for customer_id, group in frame.groupby("customer_id", sort=True)
    ]
    eb_log_scales, amount_prior_alpha, amount_prior_beta = _empirical_bayes_log_scales(
        groups
    )
    (
        eb_short_probabilities,
        short_prior_alpha,
        short_prior_beta,
    ) = _empirical_bayes_short_probabilities(groups)
    timestamps = frame["ts"]
    window_start = timestamps.min().normalize()
    window_end = timestamps.max().normalize()
    observation_days = max(1.0, float((window_end - window_start).days + 1))
    hour_counts = timestamps.dt.hour.value_counts()
    hour_weights = tuple(
        float(hour_counts.get(hour, 0) / len(frame)) for hour in range(24)
    )
    elapsed_fraction = (timestamps.dt.normalize() - window_start).dt.days / observation_days
    calendar_bins = np.minimum((elapsed_fraction * 12).astype(int), 11)
    calendar_counts = calendar_bins.value_counts()
    calendar_weights = tuple(
        float(calendar_counts.get(index, 0) / len(frame)) for index in range(12)
    )
    all_gaps = _positive_gaps_seconds(frame)
    short_gaps = all_gaps[all_gaps <= 86_400]
    if len(short_gaps) == 0:
        short_gaps = np.asarray(
            [min(86_400.0, float(np.median(all_gaps))) if len(all_gaps) else 3_600.0]
        )
    session_gap_logs = np.log(short_gaps)
    log_amounts = np.log(frame["amount"].to_numpy(dtype=float))
    global_log_center = float(np.mean(log_amounts))
    global_angles = 2 * np.pi * timestamps.dt.hour.to_numpy(dtype=float) / 24
    global_sin = float(np.mean(np.sin(global_angles)))
    global_cos = float(np.mean(np.cos(global_angles)))
    global_resultant = math.hypot(global_sin, global_cos)
    phase_cut = (math.atan2(global_sin, global_cos) + math.pi) % (2 * math.pi)
    primitive_rows: dict[str, list[list[float]]] = {
        name: [] for name in STRATUM_ORDER
    }
    for customer_id, group in groups:
        group = group.sort_values("ts", kind="stable")
        count = len(group)
        group_logs = np.log(group["amount"].to_numpy(dtype=float))
        angles = 2 * np.pi * group["ts"].dt.hour.to_numpy(dtype=float) / 24
        mean_sin = float(np.mean(np.sin(angles)))
        mean_cos = float(np.mean(np.cos(angles)))
        resultant = math.hypot(mean_sin, mean_cos)
        if resultant > 1e-12:
            phase_sin, phase_cos = mean_sin / resultant, mean_cos / resultant
        else:
            phase_norm = max(1e-12, global_resultant)
            phase_sin, phase_cos = global_sin / phase_norm, global_cos / phase_norm
        phase_latent = (math.atan2(phase_sin, phase_cos) - phase_cut) % (2 * math.pi)
        eb_resultant = min(
            0.999999,
            (count * resultant + EB_CIRCADIAN_PRIOR_STRENGTH * global_resultant)
            / (count + EB_CIRCADIAN_PRIOR_STRENGTH),
        )
        concentration = _resultant_to_concentration(eb_resultant)
        first_day = float((group["ts"].iloc[0].normalize() - window_start).days)
        last_day = float((group["ts"].iloc[-1].normalize() - window_start).days)
        base = {
            "invoice_count": float(count),
            "active_start_fraction": first_day / observation_days,
            "active_duration_fraction": (last_day - first_day) / observation_days,
            "relative_log_amount_center": float(group_logs.mean() - global_log_center),
            "eb_log_amount_scale": eb_log_scales[customer_id],
            "eb_short_session_probability": eb_short_probabilities[customer_id],
            "circadian_phase_latent": phase_latent,
            "eb_circadian_concentration": concentration,
        }
        stratum = (
            "singleton"
            if count == 1
            else "same_day_repeat"
            if first_day == last_day
            else "multiday_repeat"
        )
        primitive_rows[stratum].append(
            [base[primitive] for primitive in STRATUM_PRIMITIVES[stratum]]
        )
    arrays = {
        name: np.asarray(primitive_rows[name], dtype=float) for name in STRATUM_ORDER
    }
    for name, rows in arrays.items():
        if len(rows) < MIN_STRATUM_CUSTOMERS:
            raise ValueError(
                f"{name} has {len(rows)} customers; at least {MIN_STRATUM_CUSTOMERS} are required"
            )
    strata = []
    for name in STRATUM_ORDER:
        rows = arrays[name]
        quantiles = np.quantile(
            rows, QUANTILE_PROBABILITIES, axis=0, method="linear"
        ).T
        strata.append(
            CheckerboardStratum(
                name=name,
                customer_count=len(rows),
                primitive_quantiles=tuple(
                    tuple(map(float, values)) for values in quantiles
                ),
                cell_masses=fractional_checkerboard_masses(rows),
            )
        )
    profile = BehaviorProfileV4(
        source_dataset="UCI Online Retail 352",
        source_partition="calibration",
        source_customer_ids_sha256=customer_ids_sha256,
        source_transactions_sha256=_transactions_hash(frame),
        source_customer_count=len(groups),
        source_transaction_count=len(frame),
        quantile_probabilities=QUANTILE_PROBABILITIES,
        strata=tuple(strata),
        hour_weights=hour_weights,
        calendar_weights=calendar_weights,
        observation_days=observation_days,
        session_gap_log_mean=float(np.mean(session_gap_logs)),
        session_gap_log_std=float(np.std(session_gap_logs)),
        circadian_phase_cut=phase_cut,
        amount_variance_prior_alpha=amount_prior_alpha,
        amount_variance_prior_beta=amount_prior_beta,
        short_session_prior_alpha=short_prior_alpha,
        short_session_prior_beta=short_prior_beta,
    )
    return profile, arrays


def fit_behavior_profile_v4(
    real_calibration: Any, *, customer_ids_sha256: str
) -> BehaviorProfileV4:
    profile, _ = _fit_behavior_profile_and_rows(
        real_calibration, customer_ids_sha256=customer_ids_sha256
    )
    return profile


def novelty_check(
    profile: BehaviorProfileV4,
    calibration_primitive_rows: Mapping[str, Any],
    *,
    draws: int = NOVELTY_DRAWS,
    seed: int = SEED + 400,
    batch_size: int = 10_000,
) -> dict[str, Any]:
    """Require no exact full continuous-primitive tuple reuse."""
    if draws < 1 or batch_size < 1:
        raise ValueError("novelty draw counts must be positive")
    calibration = {
        name: {
            tuple(float(row[index]) for index, primitive in enumerate(STRATUM_PRIMITIVES[name]) if primitive != "invoice_count")
            for row in rows
        }
        for name, rows in calibration_primitive_rows.items()
    }
    exact_matches = 0
    completed = 0
    batch = 0
    while completed < draws:
        current = min(batch_size, draws - completed)
        sampled = profile.sample_primitives(current, seed=seed + batch)
        for row in sampled:
            name = (
                "singleton"
                if row["invoice_count"] == 1
                else "same_day_repeat"
                if row["active_duration_fraction"] == 0
                else "multiday_repeat"
            )
            key = tuple(
                float(row[primitive])
                for primitive in STRATUM_PRIMITIVES[name]
                if primitive != "invoice_count"
            )
            exact_matches += key in calibration[name]
        completed += current
        batch += 1
    return {
        "draws": draws,
        "exact_full_continuous_primitive_tuple_matches": exact_matches,
        "passed": exact_matches == 0,
    }


def generate_synthetic_transactions_v4(
    *, profile: BehaviorProfileV4, customers: int, seed: int
) -> Any:
    _, pd, _, _ = _dependencies()
    simulator = PopulationSimulatorV4(
        seed=seed,
        customers=customers,
        merchants_per_category=1,
        agentic_rate=0,
        behavior_profile=profile,
    )
    events = sum(
        round(customer.activity_weight) for customer in simulator.customers.values()
    )
    rows = (
        {
            "customer_id": event["txn"]["customer_id"],
            "ts": pd.Timestamp(event["ts"]).tz_convert("UTC").tz_localize(None),
            "amount": event["txn"]["amount_minor"] / 100,
        }
        for event in simulator.iter_legitimate(events)
    )
    return pd.DataFrame(rows).sort_values(["customer_id", "ts"], kind="stable").reset_index(drop=True)


def verify_retail_ii_provenance(workbook_path: Path) -> dict[str, Any]:
    manifest = verify_uci502_existing(workbook_path)
    if (
        manifest.get("dataset_id") != UCI502_DATASET_ID
        or manifest.get("source_url") != UCI502_SOURCE_URL
        or manifest.get("archive_sha256") != UCI502_ARCHIVE_SHA256
        or manifest.get("workbook_sha256") != UCI502_WORKBOOK_SHA256
        or manifest.get("license") != UCI502_LICENSE
    ):
        raise ValueError("UCI Online Retail II provenance differs from the frozen pins")
    return {
        "dataset_id": UCI502_DATASET_ID,
        "source_url": UCI502_SOURCE_URL,
        "archive_sha256": UCI502_ARCHIVE_SHA256,
        "workbook_sha256": UCI502_WORKBOOK_SHA256,
        "manifest_sha256": sha256_file(workbook_path.with_suffix(".manifest.json")),
        "license": UCI502_LICENSE,
    }


def load_retail_ii_readiness(
    workbook_path: Path, *, excluded_customer_ids: set[str]
) -> tuple[Any, dict[str, Any]]:
    """Stream only the named first worksheet and materialize only eligible rows."""
    _, pd, _, _ = _dependencies()
    provenance = verify_retail_ii_provenance(workbook_path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    selected_rows: list[dict[str, Any]] = []
    input_rows = 0
    excluded_training_rows = 0
    outside_window_rows = 0
    try:
        if not workbook.sheetnames or workbook.sheetnames[0] != UCI502_SHEET:
            raise ValueError("UCI Online Retail II first-sheet contract changed")
        values = workbook[UCI502_SHEET].iter_rows(values_only=True)
        headers = tuple(next(values))
        if headers != UCI502_HEADERS:
            raise ValueError("UCI Online Retail II column contract changed")
        customer_index = headers.index("Customer ID")
        date_index = headers.index("InvoiceDate")
        start = pd.Timestamp(UCI502_WINDOW_START)
        end = pd.Timestamp(UCI502_WINDOW_END)
        for values_row in values:
            input_rows += 1
            timestamp = pd.Timestamp(values_row[date_index])
            if pd.isna(timestamp) or timestamp < start or timestamp >= end:
                outside_window_rows += 1
                continue
            customer_id = _customer_id(values_row[customer_index])
            if customer_id is None:
                continue
            if customer_id in excluded_customer_ids:
                excluded_training_rows += 1
                continue
            source = dict(zip(headers, values_row, strict=True))
            selected_rows.append(
                {
                    "InvoiceNo": source["Invoice"],
                    "Quantity": source["Quantity"],
                    "InvoiceDate": source["InvoiceDate"],
                    "UnitPrice": source["Price"],
                    "CustomerID": source["Customer ID"],
                }
            )
    finally:
        workbook.close()
    if input_rows != UCI502_EXPECTED_ROWS:
        raise ValueError("UCI Online Retail II first-sheet row count changed")
    frame, counts = aggregate_retail_rows(selected_rows)
    return frame, {
        **provenance,
        **counts,
        "sheet": UCI502_SHEET,
        "source_input_line_items": input_rows,
        "outside_fixed_window_line_items": outside_window_rows,
        "training_customer_line_items_excluded": excluded_training_rows,
        "selected_customer_ids_sha256": _ids_hash(frame["customer_id"].unique()),
        "transaction_time_disjoint": True,
        "customer_disjoint_from_training": True,
        "customer_disjoint_from_sealed": "unknown",
        "source_independent": False,
        "second_worksheet_rows_iterated": False,
        "nonselected_transactions_materialized": False,
    }


def uci502_readiness_split(
    customer_ids: Iterable[str], *, minimum_customers: int = UCI502_MIN_SPLIT_CUSTOMERS
) -> tuple[set[str], set[str]]:
    transform: set[str] = set()
    score: set[str] = set()
    for customer_id in sorted(set(map(str, customer_ids))):
        value = int.from_bytes(
            hashlib.sha256(
                f"{UCI502_SPLIT_NAMESPACE}:{customer_id}".encode()
            ).digest()[:8],
            "big",
        )
        (transform if value < 1 << 63 else score).add(customer_id)
    if len(transform) < minimum_customers or len(score) < minimum_customers:
        raise ValueError("UCI502 readiness halves do not meet the frozen minimum size")
    return transform, score


def freeze_v4_protocol(
    reference_xlsx: Path = DEFAULT_REFERENCE,
    protocol_path: Path = DEFAULT_PROTOCOL,
) -> dict[str, Any]:
    real_calibration, reference = load_retail_partition(reference_xlsx, "calibration")
    profile, primitive_rows = _fit_behavior_profile_and_rows(
        real_calibration,
        customer_ids_sha256=_ids_hash(real_calibration["customer_id"].unique()),
    )
    novelty = novelty_check(profile, primitive_rows, draws=NOVELTY_DRAWS)
    if novelty["passed"] is not True:
        raise ValueError("v4 novelty check failed before protocol freeze")
    protocol: dict[str, Any] = {
        "schema_version": 4,
        "artifact": "entity_fidelity_v4_frozen_protocol",
        "implementation_sha256": _implementation_hashes(),
        "dependency_versions": {
            name: importlib.metadata.version(name)
            for name in (
                "numpy",
                "scipy",
                "pandas",
                "scikit-learn",
                "lightgbm",
                "openpyxl",
            )
        },
        "dependency_lock_sha256": sha256_file(ROOT / "requirements.lock"),
        "sources": {
            "training_and_sealed": {
                "dataset_id": UCI352_DATASET_ID,
                "source_url": UCI352_SOURCE_URL,
                "workbook_sha256": reference["workbook_sha256"],
                "manifest_sha256": reference["manifest_sha256"],
                "license": UCI352_LICENSE,
                "downloader_sha256": sha256_file(
                    ROOT / "scripts" / "download_retail_reference.py"
                ),
                "calibration_customer_ids_sha256": EXPECTED_PARTITION_HASHES["calibration"],
                "sealed_customer_ids_sha256": EXPECTED_PARTITION_HASHES["sealed_final"],
                "diagnostic_permitted": False,
            },
            "temporal_readiness": {
                "dataset_id": UCI502_DATASET_ID,
                "source_url": UCI502_SOURCE_URL,
                "archive_sha256": UCI502_ARCHIVE_SHA256,
                "workbook_sha256": UCI502_WORKBOOK_SHA256,
                "license": UCI502_LICENSE,
                "downloader_sha256": sha256_file(
                    ROOT / "scripts" / "download_retail_ii_reference.py"
                ),
                "sheet": UCI502_SHEET,
                "fixed_window": f"[{UCI502_WINDOW_START}, {UCI502_WINDOW_END})",
                "exclude_uci352_calibration_customers": True,
                "split_assignment": UCI502_SPLIT_ASSIGNMENT,
                "minimum_customers_per_half": UCI502_MIN_SPLIT_CUSTOMERS,
                "transaction_time_disjoint": True,
                "customer_disjoint_from_training": True,
                "customer_disjoint_from_sealed": "unknown",
                "source_independent": False,
            },
        },
        "generator": {
            "behavior_profile": profile.as_dict(),
            "behavior_profile_sha256": profile.sha256,
            "strata": {
                name: {
                    "primitives": list(STRATUM_PRIMITIVES[name]),
                    "cell_masses": STRATUM_CELL_COUNTS[name],
                }
                for name in STRATUM_ORDER
            },
            "checkerboard_bins": CHECKERBOARD_BINS,
            "quantile_probabilities": list(QUANTILE_PROBABILITIES),
            "tie_rule": CHECKERBOARD_TIE_RULE,
            "cell_allocation": CHECKERBOARD_CELL_ALLOCATION,
            "sampling": CHECKERBOARD_SAMPLING,
            "smoothing": False,
            "pseudocounts": False,
            "stored_rank_rows": False,
            "novelty": novelty,
        },
        "evaluation": {
            "seeds": list(EVALUATION_SEEDS),
            "features": list(FEATURES),
            "classifier": _classifier_protocol(),
            "exploratory_crossfit_folds": CROSS_FIT_FOLDS,
            "exploratory_assignment": EXPLORATORY_ASSIGNMENT,
            "orientation_free_statistic": "D=max(AUC,1-AUC)",
            "max_seed_statistic": True,
            "bootstrap_repeats": BOOTSTRAP_REPEATS,
            "bootstrap_quantile": 0.975,
            "shared_real_bootstrap_indices_across_seeds": True,
            "independent_synthetic_bootstrap_indices_across_seeds": True,
            "threshold_max_d_upper_95": GATE_MAX_D_UPPER_95,
        },
        "sealed_final": {
            "confirmation": f"{SEALED_FINAL_CONFIRMATION}:<protocol_sha256>",
            "exclusive_receipt_before_partition_selection": True,
            "second_invocation_fails": True,
        },
        "privacy": {
            "artifact_policy": "fixed aggregates only",
            "contains_customer_ids": False,
            "contains_raw_rows": False,
            "contains_rank_rows": False,
            "contains_evaluator_feature_rows": False,
            "contains_prediction_vectors": False,
            "differential_privacy_claimed": False,
        },
    }
    protocol["protocol_sha256"] = _protocol_hash(protocol)
    _assert_aggregate_only(protocol)
    if protocol_path.exists():
        existing = json.loads(protocol_path.read_text(encoding="utf-8"))
        if existing != protocol:
            raise FileExistsError("frozen v4 protocol exists with different content")
        return existing
    _write_exclusive_json(protocol_path, protocol)
    protocol_path.chmod(0o444)
    return protocol


def load_frozen_v4_protocol(
    protocol_path: Path = DEFAULT_PROTOCOL,
) -> dict[str, Any]:
    protocol = json.loads(protocol_path.read_text(encoding="utf-8"))
    if protocol.get("schema_version") != 4 or protocol.get("artifact") != "entity_fidelity_v4_frozen_protocol":
        raise ValueError("not a frozen v4 protocol")
    if protocol.get("protocol_sha256") != _protocol_hash(protocol):
        raise ValueError("frozen v4 protocol hash mismatch")
    if protocol.get("implementation_sha256") != _implementation_hashes():
        raise ValueError("implementation changed after the v4 protocol was frozen")
    dependency_versions = {
        name: importlib.metadata.version(name)
        for name in (
            "numpy",
            "scipy",
            "pandas",
            "scikit-learn",
            "lightgbm",
            "openpyxl",
        )
    }
    if (
        protocol.get("dependency_versions") != dependency_versions
        or protocol.get("dependency_lock_sha256")
        != sha256_file(ROOT / "requirements.lock")
    ):
        raise ValueError("v4 dependency bindings changed")
    training_source = protocol.get("sources", {}).get("training_and_sealed", {})
    source = protocol.get("sources", {}).get("temporal_readiness", {})
    if (
        training_source.get("dataset_id") != UCI352_DATASET_ID
        or training_source.get("source_url") != UCI352_SOURCE_URL
        or training_source.get("license") != UCI352_LICENSE
        or training_source.get("workbook_sha256") != UCI352_WORKBOOK_SHA256
        or training_source.get("calibration_customer_ids_sha256")
        != EXPECTED_PARTITION_HASHES["calibration"]
        or training_source.get("sealed_customer_ids_sha256")
        != EXPECTED_PARTITION_HASHES["sealed_final"]
        or training_source.get("diagnostic_permitted") is not False
        or training_source.get("downloader_sha256")
        != sha256_file(ROOT / "scripts" / "download_retail_reference.py")
        or source.get("dataset_id") != UCI502_DATASET_ID
        or source.get("source_url") != UCI502_SOURCE_URL
        or source.get("license") != UCI502_LICENSE
        or source.get("archive_sha256") != UCI502_ARCHIVE_SHA256
        or source.get("workbook_sha256") != UCI502_WORKBOOK_SHA256
        or source.get("sheet") != UCI502_SHEET
        or source.get("fixed_window")
        != f"[{UCI502_WINDOW_START}, {UCI502_WINDOW_END})"
        or source.get("exclude_uci352_calibration_customers") is not True
        or source.get("split_assignment") != UCI502_SPLIT_ASSIGNMENT
        or source.get("minimum_customers_per_half") != UCI502_MIN_SPLIT_CUSTOMERS
        or source.get("transaction_time_disjoint") is not True
        or source.get("customer_disjoint_from_training") is not True
        or source.get("customer_disjoint_from_sealed") != "unknown"
        or source.get("source_independent") is not False
        or source.get("downloader_sha256")
        != sha256_file(ROOT / "scripts" / "download_retail_ii_reference.py")
    ):
        raise ValueError("v4 readiness source pins changed")
    profile = BehaviorProfileV4.from_dict(protocol["generator"]["behavior_profile"])
    if (
        profile.sha256 != protocol["generator"].get("behavior_profile_sha256")
        or profile.source_dataset != "UCI Online Retail 352"
        or profile.source_customer_ids_sha256
        != EXPECTED_PARTITION_HASHES["calibration"]
    ):
        raise ValueError("frozen v4 behavior profile hash mismatch")
    generator = protocol["generator"]
    expected_strata = {
        name: {
            "primitives": list(STRATUM_PRIMITIVES[name]),
            "cell_masses": STRATUM_CELL_COUNTS[name],
        }
        for name in STRATUM_ORDER
    }
    if (
        generator.get("strata") != expected_strata
        or generator.get("checkerboard_bins") != CHECKERBOARD_BINS
        or tuple(generator.get("quantile_probabilities", ()))
        != QUANTILE_PROBABILITIES
        or generator.get("tie_rule") != CHECKERBOARD_TIE_RULE
        or generator.get("cell_allocation") != CHECKERBOARD_CELL_ALLOCATION
        or generator.get("sampling") != CHECKERBOARD_SAMPLING
        or generator.get("smoothing") is not False
        or generator.get("pseudocounts") is not False
        or generator.get("stored_rank_rows") is not False
    ):
        raise ValueError("v4 frozen generator contract changed")
    novelty = generator.get("novelty", {})
    if novelty.get("draws") != NOVELTY_DRAWS or novelty.get("passed") is not True or novelty.get("exact_full_continuous_primitive_tuple_matches") != 0:
        raise ValueError("frozen v4 novelty gate is absent or failed")
    evaluation = protocol.get("evaluation", {})
    if (
        tuple(evaluation.get("seeds", ())) != EVALUATION_SEEDS
        or tuple(evaluation.get("features", ())) != FEATURES
        or evaluation.get("classifier") != _classifier_protocol()
        or evaluation.get("exploratory_crossfit_folds") != CROSS_FIT_FOLDS
        or evaluation.get("exploratory_assignment") != EXPLORATORY_ASSIGNMENT
        or evaluation.get("orientation_free_statistic") != "D=max(AUC,1-AUC)"
        or evaluation.get("max_seed_statistic") is not True
        or evaluation.get("bootstrap_repeats") != BOOTSTRAP_REPEATS
        or evaluation.get("bootstrap_quantile") != 0.975
        or evaluation.get("shared_real_bootstrap_indices_across_seeds") is not True
        or evaluation.get("independent_synthetic_bootstrap_indices_across_seeds")
        is not True
        or evaluation.get("threshold_max_d_upper_95") != GATE_MAX_D_UPPER_95
    ):
        raise ValueError("v4 evaluation contract changed")
    expected_privacy = {
        "artifact_policy": "fixed aggregates only",
        "contains_customer_ids": False,
        "contains_raw_rows": False,
        "contains_rank_rows": False,
        "contains_evaluator_feature_rows": False,
        "contains_prediction_vectors": False,
        "differential_privacy_claimed": False,
    }
    if protocol.get("privacy") != expected_privacy:
        raise ValueError("v4 privacy contract changed")
    if protocol.get("sealed_final") != {
        "confirmation": f"{SEALED_FINAL_CONFIRMATION}:<protocol_sha256>",
        "exclusive_receipt_before_partition_selection": True,
        "second_invocation_fails": True,
    }:
        raise ValueError("v4 sealed-final contract changed")
    _assert_aggregate_only(protocol)
    return protocol


def _synthetic_group(
    profile: BehaviorProfileV4, *, customers: int, seed: int, prefix: str
) -> Any:
    frame = generate_synthetic_transactions_v4(
        profile=profile, customers=customers, seed=seed
    )
    frame["customer_id"] = prefix + frame["customer_id"].astype(str)
    return frame


def _crossfit_seed(
    real_calibration: Any, *, seed: int, fold_numbers: Iterable[int] | None = None
) -> list[dict[str, Any]]:
    np, pd, _, roc_auc_score = _dependencies()
    customer_ids = sorted(set(real_calibration["customer_id"]))
    fold_by_id = {
        customer_id: int.from_bytes(
            hashlib.sha256(
                f"{seed}:uci-v4-crossfit:{customer_id}".encode()
            ).digest()[:8],
            "big",
        )
        % CROSS_FIT_FOLDS
        for customer_id in customer_ids
    }
    folds = []
    for fold in fold_numbers if fold_numbers is not None else range(CROSS_FIT_FOLDS):
        evaluation_ids = {
            customer_id for customer_id, assigned in fold_by_id.items() if assigned == fold
        }
        training_ids = set(customer_ids) - evaluation_ids
        if not training_ids or not evaluation_ids:
            raise ValueError("v4 exploratory cross-fit fold is empty")
        real_training = real_calibration.loc[
            real_calibration["customer_id"].isin(training_ids)
        ].copy()
        real_evaluation = real_calibration.loc[
            real_calibration["customer_id"].isin(evaluation_ids)
        ].copy()
        profile = fit_behavior_profile_v4(
            real_training, customer_ids_sha256=_ids_hash(training_ids)
        )
        synthetic_training = _synthetic_group(
            profile,
            customers=len(training_ids),
            seed=seed + 10_000 * fold + 101,
            prefix=f"v4-train-{seed}-{fold}-",
        )
        synthetic_evaluation = _synthetic_group(
            profile,
            customers=len(evaluation_ids),
            seed=seed + 10_000 * fold + 202,
            prefix=f"v4-eval-{seed}-{fold}-",
        )
        real_transform, synthetic_transform, _ = fit_calibrations(
            real_training, synthetic_training
        )
        real_training_features = customer_features(real_training, real_transform)
        synthetic_training_features = customer_features(
            synthetic_training, synthetic_transform
        )
        classifier = _classifier(seed + fold)
        classifier.fit(
            pd.concat(
                (real_training_features, synthetic_training_features),
                ignore_index=True,
            ),
            np.concatenate(
                (
                    np.zeros(len(real_training_features), dtype="int8"),
                    np.ones(len(synthetic_training_features), dtype="int8"),
                )
            ),
        )
        real_evaluation_features = customer_features(real_evaluation, real_transform)
        real_predictions = classifier.predict_proba(real_evaluation_features)[:, 1]
        synthetic_predictions = classifier.predict_proba(
            customer_features(synthetic_evaluation, synthetic_transform)
        )[:, 1]
        labels = np.concatenate(
            (np.zeros(len(real_predictions)), np.ones(len(synthetic_predictions)))
        )
        auc = float(
            roc_auc_score(
                labels, np.concatenate((real_predictions, synthetic_predictions))
            )
        )
        folds.append(
            {
                "fold": fold,
                "training_real_customers": len(training_ids),
                "evaluation_real_customers": len(evaluation_ids),
                "evaluation_synthetic_customers": len(synthetic_predictions),
                "training_customer_ids_sha256": _ids_hash(training_ids),
                "evaluation_customer_ids_sha256": _ids_hash(evaluation_ids),
                "profile_sha256": profile.sha256,
                "auc": auc,
                "d": max(auc, 1 - auc),
                "_real_customer_ids": tuple(real_evaluation_features.index),
                "_real_predictions": real_predictions,
                "_synthetic_predictions": synthetic_predictions,
            }
        )
    return folds


def _crossfit_max_seed_bootstrap(
    seed_folds: Mapping[int, list[dict[str, Any]]], *, repeats: int, seed: int
) -> tuple[list[float], float, float]:
    """Share bootstrapped real customer identities across cross-fit seeds."""
    np, _, _, roc_auc_score = _dependencies()
    if repeats < 100:
        raise ValueError("at least 100 stratified customer bootstraps are required")
    first_folds = next(iter(seed_folds.values()))
    customer_ids = sorted(
        customer_id
        for fold in first_folds
        for customer_id in fold["_real_customer_ids"]
    )
    lookups: dict[int, dict[str, tuple[int, float]]] = {}
    synthetic_by_fold: dict[int, dict[int, Any]] = {}
    for evaluation_seed, folds in seed_folds.items():
        lookup: dict[str, tuple[int, float]] = {}
        synthetic_by_fold[evaluation_seed] = {}
        for fold in folds:
            synthetic_by_fold[evaluation_seed][fold["fold"]] = fold[
                "_synthetic_predictions"
            ]
            for customer_id, prediction in zip(
                fold["_real_customer_ids"], fold["_real_predictions"], strict=True
            ):
                lookup[str(customer_id)] = (fold["fold"], float(prediction))
        if set(lookup) != set(customer_ids):
            raise ValueError("cross-fit seeds must cover the same real customers")
        lookups[evaluation_seed] = lookup
    rng = np.random.default_rng(seed)
    maxima = []
    for _ in range(repeats):
        sampled_ids = [
            customer_ids[index]
            for index in rng.integers(0, len(customer_ids), len(customer_ids))
        ]
        seed_values = []
        for evaluation_seed, lookup in lookups.items():
            real_by_fold: dict[int, list[float]] = {
                fold: [] for fold in range(CROSS_FIT_FOLDS)
            }
            for customer_id in sampled_ids:
                fold, prediction = lookup[customer_id]
                real_by_fold[fold].append(prediction)
            weighted = 0.0
            total = 0
            for fold, real_values in real_by_fold.items():
                if not real_values:
                    continue
                real = np.asarray(real_values, dtype=float)
                synthetic_source = synthetic_by_fold[evaluation_seed][fold]
                synthetic = synthetic_source[
                    rng.integers(0, len(synthetic_source), len(real))
                ]
                labels = np.concatenate((np.zeros(len(real)), np.ones(len(synthetic))))
                auc = float(
                    roc_auc_score(labels, np.concatenate((real, synthetic)))
                )
                weighted += len(real) * max(auc, 1 - auc)
                total += len(real)
            seed_values.append(weighted / total)
        maxima.append(max(seed_values))
    low, high = np.quantile(maxima, (0.025, 0.975))
    return maxima, float(low), float(high)


def run_calibration_exploratory(
    reference_xlsx: Path = DEFAULT_REFERENCE,
    output: Path = DEFAULT_EXPLORATORY,
    *,
    protocol_path: Path = DEFAULT_PROTOCOL,
) -> dict[str, Any]:
    np, _, _, _ = _dependencies()
    protocol = load_frozen_v4_protocol(protocol_path)
    real_calibration, reference = load_retail_partition(reference_xlsx, "calibration")
    profile = fit_behavior_profile_v4(
        real_calibration,
        customer_ids_sha256=_ids_hash(real_calibration["customer_id"].unique()),
    )
    if profile.sha256 != protocol["generator"]["behavior_profile_sha256"]:
        raise ValueError("all-calibration v4 profile differs from the frozen protocol")
    seed_folds = {
        seed: _crossfit_seed(real_calibration, seed=seed) for seed in EVALUATION_SEEDS
    }
    seed_point = {
        str(seed): sum(
            fold["d"] * fold["evaluation_real_customers"] for fold in folds
        )
        / sum(fold["evaluation_real_customers"] for fold in folds)
        for seed, folds in seed_folds.items()
    }
    bootstrap, ci_low, ci_high = _crossfit_max_seed_bootstrap(
        seed_folds, repeats=BOOTSTRAP_REPEATS, seed=SEED + 30
    )
    result = {
        "schema_version": 4,
        "artifact": "entity_fidelity_v4_calibration_exploratory",
        "protocol_sha256": protocol["protocol_sha256"],
        "calibration_only": True,
        "diagnostic_transactions_accessed": False,
        "uci502_transactions_accessed": False,
        "sealed_final_transactions_accessed": False,
        "reference": {
            "workbook_sha256": reference["workbook_sha256"],
            "partition": "calibration",
            "customer_ids_sha256": EXPECTED_PARTITION_HASHES["calibration"],
            "customers": int(real_calibration["customer_id"].nunique()),
            "transactions": len(real_calibration),
        },
        "seeds": {
            str(seed): {
                "fold_weighted_d": seed_point[str(seed)],
                "folds": [
                    {key: value for key, value in fold.items() if not key.startswith("_")}
                    for fold in folds
                ],
            }
            for seed, folds in seed_folds.items()
        },
        "metrics": {
            "orientation_free_max_seed_d": max(seed_point.values()),
            "bootstrap_max_seed_d_mean": float(np.mean(bootstrap)),
            "bootstrap_max_seed_d_ci_95": [ci_low, ci_high],
            "threshold_max_d_upper_95": GATE_MAX_D_UPPER_95,
        },
        "gate_passed": ci_high <= GATE_MAX_D_UPPER_95,
    }
    _write_artifact(output, result)
    return result


def _load_passed_exploratory(path: Path, protocol: Mapping[str, Any]) -> dict[str, Any]:
    result = json.loads(path.read_text(encoding="utf-8"))
    if (
        result.get("schema_version") != 4
        or result.get("artifact") != "entity_fidelity_v4_calibration_exploratory"
        or result.get("protocol_sha256") != protocol["protocol_sha256"]
        or not _gate_metrics_valid(result)
        or result.get("calibration_only") is not True
        or result.get("diagnostic_transactions_accessed") is not False
        or result.get("uci502_transactions_accessed") is not False
        or result.get("sealed_final_transactions_accessed") is not False
    ):
        raise ValueError("v4 exploratory gate is absent, failed, or not protocol-bound")
    return result


def _readiness_seed_predictions(
    real_calibration: Any,
    historical_transform: Any,
    historical_score: Any,
    profile: BehaviorProfileV4,
    *,
    seed: int,
) -> dict[str, Any]:
    np, pd, _, roc_auc_score = _dependencies()
    training_customers = int(real_calibration["customer_id"].nunique())
    evaluation_customers = int(historical_score["customer_id"].nunique())
    synthetic_training = _synthetic_group(
        profile,
        customers=training_customers,
        seed=seed + 500_001,
        prefix=f"readiness-train-{seed}-",
    )
    synthetic_evaluation = _synthetic_group(
        profile,
        customers=evaluation_customers,
        seed=seed + 500_002,
        prefix=f"readiness-eval-{seed}-",
    )
    real_transform, synthetic_transform, _ = fit_calibrations(
        real_calibration, synthetic_training
    )
    real_training_features = customer_features(real_calibration, real_transform)
    synthetic_training_features = customer_features(
        synthetic_training, synthetic_transform
    )
    classifier = _classifier(seed)
    classifier.fit(
        pd.concat((real_training_features, synthetic_training_features), ignore_index=True),
        np.concatenate(
            (
                np.zeros(len(real_training_features), dtype="int8"),
                np.ones(len(synthetic_training_features), dtype="int8"),
            )
        ),
    )
    real_predictions = classifier.predict_proba(
        customer_features(historical_score, historical_transform)
    )[:, 1]
    synthetic_predictions = classifier.predict_proba(
        customer_features(synthetic_evaluation, synthetic_transform)
    )[:, 1]
    labels = np.concatenate(
        (np.zeros(len(real_predictions)), np.ones(len(synthetic_predictions)))
    )
    auc = float(
        roc_auc_score(labels, np.concatenate((real_predictions, synthetic_predictions)))
    )
    return {
        "auc": auc,
        "d": max(auc, 1 - auc),
        "real_customers": len(real_predictions),
        "synthetic_customers": len(synthetic_predictions),
        "_real_predictions": real_predictions,
        "_synthetic_predictions": synthetic_predictions,
    }


def _max_seed_bootstrap(
    seed_predictions: Mapping[int, dict[str, Any]], *, repeats: int, seed: int
) -> tuple[list[float], float, float]:
    np, _, _, roc_auc_score = _dependencies()
    if repeats < 100:
        raise ValueError("at least 100 stratified customer bootstraps are required")
    rng = np.random.default_rng(seed)
    first = next(iter(seed_predictions.values()))
    real_count = len(first["_real_predictions"])
    if any(
        len(value["_real_predictions"]) != real_count
        for value in seed_predictions.values()
    ):
        raise ValueError("shared real bootstrap requires the same evaluation customers")
    maxima = []
    for _ in range(repeats):
        real_indices = rng.integers(0, real_count, real_count)
        values = []
        for result in seed_predictions.values():
            real = result["_real_predictions"][real_indices]
            synthetic_source = result["_synthetic_predictions"]
            synthetic = synthetic_source[
                rng.integers(0, len(synthetic_source), len(synthetic_source))
            ]
            labels = np.concatenate((np.zeros(len(real)), np.ones(len(synthetic))))
            auc = float(roc_auc_score(labels, np.concatenate((real, synthetic))))
            values.append(max(auc, 1 - auc))
        maxima.append(max(values))
    low, high = np.quantile(maxima, (0.025, 0.975))
    return maxima, float(low), float(high)


def run_temporal_readiness(
    reference_xlsx: Path = DEFAULT_REFERENCE,
    readiness_xlsx: Path = DEFAULT_READINESS_REFERENCE,
    output: Path = DEFAULT_READINESS,
    *,
    protocol_path: Path = DEFAULT_PROTOCOL,
    exploratory_path: Path = DEFAULT_EXPLORATORY,
) -> dict[str, Any]:
    np, pd, _, _ = _dependencies()
    protocol = load_frozen_v4_protocol(protocol_path)
    _load_passed_exploratory(exploratory_path, protocol)
    real_calibration, _ = load_retail_partition(reference_xlsx, "calibration")
    calibration_ids = set(real_calibration["customer_id"].astype(str))
    profile = fit_behavior_profile_v4(
        real_calibration, customer_ids_sha256=_ids_hash(calibration_ids)
    )
    if profile.sha256 != protocol["generator"]["behavior_profile_sha256"]:
        raise ValueError("all-calibration v4 profile differs from the frozen protocol")
    historical, provenance = load_retail_ii_readiness(
        readiness_xlsx, excluded_customer_ids=calibration_ids
    )
    if calibration_ids & set(historical["customer_id"].astype(str)):
        raise ValueError("UCI502 readiness contains a UCI352 calibration customer")
    transform_ids, score_ids = uci502_readiness_split(historical["customer_id"].unique())
    transform = historical.loc[historical["customer_id"].isin(transform_ids)].copy()
    score = historical.loc[historical["customer_id"].isin(score_ids)].copy()
    historical_transform = DomainCalibration(
        amount_log_reference=np.sort(
            np.log1p(transform["amount"].to_numpy(dtype=float))
        ),
        window_start=pd.Timestamp(UCI502_WINDOW_START),
        window_end=pd.Timestamp(UCI502_WINDOW_END),
        time_scale=1.0,
    )
    seed_predictions = {
        seed: _readiness_seed_predictions(
            real_calibration,
            historical_transform,
            score,
            profile,
            seed=seed,
        )
        for seed in EVALUATION_SEEDS
    }
    bootstrap, ci_low, ci_high = _max_seed_bootstrap(
        seed_predictions, repeats=BOOTSTRAP_REPEATS, seed=SEED + 31
    )
    result = {
        "schema_version": 4,
        "artifact": "entity_fidelity_v4_temporal_readiness",
        "protocol_sha256": protocol["protocol_sha256"],
        "exploratory_artifact_sha256": sha256_file(exploratory_path),
        "diagnostic_transactions_accessed": False,
        "sealed_final_transactions_accessed": False,
        "source": {
            "dataset_id": UCI502_DATASET_ID,
            "workbook_sha256": provenance["workbook_sha256"],
            "manifest_sha256": provenance["manifest_sha256"],
            "sheet": UCI502_SHEET,
            "fixed_window": f"[{UCI502_WINDOW_START}, {UCI502_WINDOW_END})",
            "transaction_time_disjoint": True,
            "customer_disjoint_from_training": True,
            "customer_disjoint_from_sealed": "unknown",
            "source_independent": False,
            "second_worksheet_rows_iterated": False,
            "transform_customers": len(transform_ids),
            "score_customers": len(score_ids),
            "transform_customer_ids_sha256": _ids_hash(transform_ids),
            "score_customer_ids_sha256": _ids_hash(score_ids),
            "historical_amount_ecdf_fit_on_transform_only": True,
            "transform_used_for_generator_or_classifier_training": False,
        },
        "seeds": {
            str(seed): {
                key: value for key, value in prediction.items() if not key.startswith("_")
            }
            for seed, prediction in seed_predictions.items()
        },
        "metrics": {
            "orientation_free_max_seed_d": max(
                prediction["d"] for prediction in seed_predictions.values()
            ),
            "bootstrap_max_seed_d_mean": float(np.mean(bootstrap)),
            "bootstrap_max_seed_d_ci_95": [ci_low, ci_high],
            "threshold_max_d_upper_95": GATE_MAX_D_UPPER_95,
        },
        "gate_passed": ci_high <= GATE_MAX_D_UPPER_95,
    }
    _write_artifact(output, result)
    return result


def _load_passed_readiness(path: Path, protocol: Mapping[str, Any]) -> dict[str, Any]:
    result = json.loads(path.read_text(encoding="utf-8"))
    source = result.get("source", {})
    if (
        result.get("schema_version") != 4
        or result.get("artifact") != "entity_fidelity_v4_temporal_readiness"
        or result.get("protocol_sha256") != protocol["protocol_sha256"]
        or not _gate_metrics_valid(result)
        or result.get("diagnostic_transactions_accessed") is not False
        or result.get("sealed_final_transactions_accessed") is not False
        or source.get("source_independent") is not False
        or source.get("customer_disjoint_from_training") is not True
        or source.get("transaction_time_disjoint") is not True
        or source.get("customer_disjoint_from_sealed") != "unknown"
        or source.get("second_worksheet_rows_iterated") is not False
        or source.get("workbook_sha256") != UCI502_WORKBOOK_SHA256
        or source.get("sheet") != UCI502_SHEET
        or source.get("fixed_window")
        != f"[{UCI502_WINDOW_START}, {UCI502_WINDOW_END})"
        or source.get("transform_customers", 0) < UCI502_MIN_SPLIT_CUSTOMERS
        or source.get("score_customers", 0) < UCI502_MIN_SPLIT_CUSTOMERS
        or source.get("historical_amount_ecdf_fit_on_transform_only") is not True
        or source.get("transform_used_for_generator_or_classifier_training")
        is not False
    ):
        raise ValueError("v4 temporal readiness is absent, failed, or not protocol-bound")
    return result


def _sealed_seed_predictions(
    real_calibration: Any,
    real_sealed: Any,
    profile: BehaviorProfileV4,
    *,
    seed: int,
) -> dict[str, Any]:
    np, pd, _, roc_auc_score = _dependencies()
    training_customers = int(real_calibration["customer_id"].nunique())
    evaluation_customers = int(real_sealed["customer_id"].nunique())
    synthetic_training = _synthetic_group(
        profile,
        customers=training_customers,
        seed=seed + 700_001,
        prefix=f"sealed-train-{seed}-",
    )
    synthetic_evaluation = _synthetic_group(
        profile,
        customers=evaluation_customers,
        seed=seed + 700_002,
        prefix=f"sealed-eval-{seed}-",
    )
    real_transform, synthetic_transform, _ = fit_calibrations(
        real_calibration, synthetic_training
    )
    real_training_features = customer_features(real_calibration, real_transform)
    synthetic_training_features = customer_features(
        synthetic_training, synthetic_transform
    )
    classifier = _classifier(seed)
    classifier.fit(
        pd.concat((real_training_features, synthetic_training_features), ignore_index=True),
        np.concatenate(
            (
                np.zeros(len(real_training_features), dtype="int8"),
                np.ones(len(synthetic_training_features), dtype="int8"),
            )
        ),
    )
    real_predictions = classifier.predict_proba(
        customer_features(real_sealed, real_transform)
    )[:, 1]
    synthetic_predictions = classifier.predict_proba(
        customer_features(synthetic_evaluation, synthetic_transform)
    )[:, 1]
    labels = np.concatenate(
        (np.zeros(len(real_predictions)), np.ones(len(synthetic_predictions)))
    )
    auc = float(
        roc_auc_score(labels, np.concatenate((real_predictions, synthetic_predictions)))
    )
    return {
        "auc": auc,
        "d": max(auc, 1 - auc),
        "real_customers": len(real_predictions),
        "synthetic_customers": len(synthetic_predictions),
        "_real_predictions": real_predictions,
        "_synthetic_predictions": synthetic_predictions,
    }


def run_entity_fidelity_v4_final(
    reference_xlsx: Path = DEFAULT_REFERENCE,
    output: Path = DEFAULT_OUTPUT,
    *,
    protocol_path: Path = DEFAULT_PROTOCOL,
    readiness_path: Path = DEFAULT_READINESS,
    receipt_path: Path = DEFAULT_SEAL_RECEIPT,
    sealed_final_confirmation: str | None = None,
) -> dict[str, Any]:
    np, _, _, _ = _dependencies()
    protocol = load_frozen_v4_protocol(protocol_path)
    expected_confirmation = f"{SEALED_FINAL_CONFIRMATION}:{protocol['protocol_sha256']}"
    if sealed_final_confirmation != expected_confirmation:
        raise ValueError("sealed-final scoring requires protocol-hash-bound one-time approval")
    _load_passed_readiness(readiness_path, protocol)
    if output.exists():
        raise FileExistsError("final entity-fidelity output already exists")
    if receipt_path.exists():
        raise FileExistsError("v4 sealed-final one-shot receipt already exists")
    real_calibration, _ = load_retail_partition(reference_xlsx, "calibration")
    profile = fit_behavior_profile_v4(
        real_calibration,
        customer_ids_sha256=_ids_hash(real_calibration["customer_id"].unique()),
    )
    if profile.sha256 != protocol["generator"]["behavior_profile_sha256"]:
        raise ValueError("all-calibration v4 profile differs from the frozen protocol")
    receipt = {
        "schema_version": 4,
        "artifact": "entity_fidelity_v4_sealed_final_one_shot_receipt",
        "protocol_sha256": protocol["protocol_sha256"],
        "status": "consumed",
        "confirmation_sha256": hashlib.sha256(
            sealed_final_confirmation.encode()
        ).hexdigest(),
        "sealed_partition_customer_ids_sha256": EXPECTED_PARTITION_HASHES["sealed_final"],
        "readiness_artifact_sha256": sha256_file(readiness_path),
    }
    receipt["receipt_sha256"] = _canonical_hash(receipt)
    _write_artifact(receipt_path, receipt)
    real_sealed, sealed_reference = load_retail_partition(reference_xlsx, "sealed_final")
    seed_predictions = {
        seed: _sealed_seed_predictions(real_calibration, real_sealed, profile, seed=seed)
        for seed in EVALUATION_SEEDS
    }
    bootstrap, ci_low, ci_high = _final_bootstrap(
        seed_predictions, repeats=BOOTSTRAP_REPEATS, seed=SEED + 32
    )
    result = {
        "schema_version": 4,
        "artifact": "entity_fidelity_v4_sealed_final",
        "protocol_sha256": protocol["protocol_sha256"],
        "protocol": {
            "protocol_sha256": protocol["protocol_sha256"],
            "evaluation_partition": "sealed_final",
            "sealed_final_scored": True,
            "readiness_gate_passed": True,
            "orientation_free_statistic": "D=max(AUC,1-AUC)",
            "bootstrap_repeats": BOOTSTRAP_REPEATS,
            "shared_real_bootstrap_indices_across_seeds": True,
            "independent_synthetic_bootstrap_indices_across_seeds": True,
        },
        "bindings": {
            "readiness_artifact_sha256": sha256_file(readiness_path),
            "sealed_receipt_sha256": receipt["receipt_sha256"],
            "behavior_profile_sha256": profile.sha256,
        },
        "behavior_profile": profile.as_dict(),
        "reference": {
            "workbook_sha256": sealed_reference["workbook_sha256"],
            "calibration_customer_ids_sha256": EXPECTED_PARTITION_HASHES["calibration"],
            "sealed_final_customer_ids_sha256": EXPECTED_PARTITION_HASHES["sealed_final"],
            "diagnostic_transactions_accessed": False,
            "calibration_customers": int(real_calibration["customer_id"].nunique()),
            "sealed_final_customers": int(real_sealed["customer_id"].nunique()),
        },
        "seeds": {
            str(seed): {
                key: value for key, value in prediction.items() if not key.startswith("_")
            }
            for seed, prediction in seed_predictions.items()
        },
        "metrics": {
            "orientation_free_max_seed_d": max(
                prediction["d"] for prediction in seed_predictions.values()
            ),
            "bootstrap_max_seed_d_mean": float(np.mean(bootstrap)),
            "bootstrap_max_seed_d_ci_95": [ci_low, ci_high],
            "threshold_max_d_upper_95": GATE_MAX_D_UPPER_95,
            "gate_passed": ci_high <= GATE_MAX_D_UPPER_95,
        },
    }
    _write_artifact(output, result)
    return result


def verify_v4_artifacts(
    *,
    protocol_path: Path = DEFAULT_PROTOCOL,
    exploratory_path: Path = DEFAULT_EXPLORATORY,
    readiness_path: Path = DEFAULT_READINESS,
    receipt_path: Path = DEFAULT_SEAL_RECEIPT,
    final_path: Path = DEFAULT_OUTPUT,
) -> dict[str, Any]:
    protocol = load_frozen_v4_protocol(protocol_path)
    checks = {
        "exploratory_exists": exploratory_path.exists(),
        "readiness_exists": readiness_path.exists(),
        "sealed_artifacts_complete": receipt_path.exists() and final_path.exists(),
    }
    if exploratory_path.exists():
        exploratory = json.loads(exploratory_path.read_text(encoding="utf-8"))
        checks.update(
            {
                "exploratory_protocol_bound": exploratory.get("protocol_sha256")
                == protocol["protocol_sha256"],
                "exploratory_gate_passed": _gate_metrics_valid(exploratory),
                "exploratory_excluded_sensitive_sources": exploratory.get(
                    "diagnostic_transactions_accessed"
                )
                is False
                and exploratory.get("uci502_transactions_accessed") is False
                and exploratory.get("sealed_final_transactions_accessed") is False,
            }
        )
    if readiness_path.exists():
        readiness = json.loads(readiness_path.read_text(encoding="utf-8"))
        source = readiness.get("source", {})
        checks.update(
            {
                "readiness_protocol_bound": readiness.get("protocol_sha256")
                == protocol["protocol_sha256"],
                "readiness_exploratory_bound": exploratory_path.exists()
                and readiness.get("exploratory_artifact_sha256")
                == sha256_file(exploratory_path),
                "readiness_gate_passed": _gate_metrics_valid(readiness),
                "readiness_source_honest": source.get("source_independent") is False
                and source.get("customer_disjoint_from_sealed") == "unknown"
                and source.get("customer_disjoint_from_training") is True
                and source.get("transaction_time_disjoint") is True
                and source.get("second_worksheet_rows_iterated") is False,
            }
        )
    if receipt_path.exists() and final_path.exists():
        receipt = json.loads(receipt_path.read_text(encoding="utf-8"))
        final = json.loads(final_path.read_text(encoding="utf-8"))
        expected_receipt_hash = _canonical_hash(
            {key: value for key, value in receipt.items() if key != "receipt_sha256"}
        )
        checks.update(
            {
                "receipt_protocol_bound": receipt.get("protocol_sha256")
                == protocol["protocol_sha256"],
                "receipt_consumed": receipt.get("status") == "consumed",
                "receipt_hash_valid": receipt.get("receipt_sha256")
                == expected_receipt_hash,
                "receipt_readiness_bound": readiness_path.exists()
                and receipt.get("readiness_artifact_sha256")
                == sha256_file(readiness_path),
                "final_protocol_bound": final.get("protocol", {}).get(
                    "protocol_sha256"
                )
                == protocol["protocol_sha256"],
                "final_is_sealed": final.get("protocol", {}).get(
                    "evaluation_partition"
                )
                == "sealed_final"
                and final.get("protocol", {}).get("sealed_final_scored") is True,
                "final_declares_readiness_passed": final.get("protocol", {}).get(
                    "readiness_gate_passed"
                )
                is True,
                "final_readiness_bound": final.get("bindings", {}).get(
                    "readiness_artifact_sha256"
                )
                == (sha256_file(readiness_path) if readiness_path.exists() else None),
                "final_receipt_bound": final.get("bindings", {}).get(
                    "sealed_receipt_sha256"
                )
                == receipt.get("receipt_sha256"),
                "final_profile_bound": _canonical_hash(final.get("behavior_profile", {}))
                == protocol["generator"]["behavior_profile_sha256"],
                "final_excluded_diagnostic": final.get("reference", {}).get(
                    "diagnostic_transactions_accessed"
                )
                is False,
                "final_gate_passed": _gate_metrics_valid(final),
            }
        )
    return {
        "protocol_sha256": protocol["protocol_sha256"],
        "checks": checks,
        "verified": all(checks.values()),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    freeze = commands.add_parser("freeze", help="freeze the calibration-fitted v4 protocol")
    freeze.add_argument("--reference-xlsx", type=Path, default=DEFAULT_REFERENCE)
    freeze.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    exploratory = commands.add_parser(
        "exploratory", help="run frozen UCI352 calibration cross-fit futility"
    )
    exploratory.add_argument("--reference-xlsx", type=Path, default=DEFAULT_REFERENCE)
    exploratory.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    exploratory.add_argument("--output", type=Path, default=DEFAULT_EXPLORATORY)
    readiness = commands.add_parser(
        "readiness", help="run UCI502 temporal readiness after exploratory PASS"
    )
    readiness.add_argument("--reference-xlsx", type=Path, default=DEFAULT_REFERENCE)
    readiness.add_argument(
        "--readiness-xlsx", type=Path, default=DEFAULT_READINESS_REFERENCE
    )
    readiness.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    readiness.add_argument("--exploratory", type=Path, default=DEFAULT_EXPLORATORY)
    readiness.add_argument("--output", type=Path, default=DEFAULT_READINESS)
    sealed = commands.add_parser(
        "sealed", help="consume the one-shot receipt and score sealed_final once"
    )
    sealed.add_argument("--reference-xlsx", type=Path, default=DEFAULT_REFERENCE)
    sealed.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    sealed.add_argument("--readiness", type=Path, default=DEFAULT_READINESS)
    sealed.add_argument("--receipt", type=Path, default=DEFAULT_SEAL_RECEIPT)
    sealed.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    sealed.add_argument("--confirmation", required=True)
    verify = commands.add_parser("verify", help="verify all v4 artifact bindings")
    verify.add_argument("--protocol", type=Path, default=DEFAULT_PROTOCOL)
    verify.add_argument("--exploratory", type=Path, default=DEFAULT_EXPLORATORY)
    verify.add_argument("--readiness", type=Path, default=DEFAULT_READINESS)
    verify.add_argument("--receipt", type=Path, default=DEFAULT_SEAL_RECEIPT)
    verify.add_argument("--final", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if args.command == "freeze":
        result = freeze_v4_protocol(args.reference_xlsx, args.protocol)
    elif args.command == "exploratory":
        result = run_calibration_exploratory(
            args.reference_xlsx, args.output, protocol_path=args.protocol
        )
    elif args.command == "readiness":
        result = run_temporal_readiness(
            args.reference_xlsx,
            args.readiness_xlsx,
            args.output,
            protocol_path=args.protocol,
            exploratory_path=args.exploratory,
        )
    elif args.command == "sealed":
        result = run_entity_fidelity_v4_final(
            args.reference_xlsx,
            args.output,
            protocol_path=args.protocol,
            readiness_path=args.readiness,
            receipt_path=args.receipt,
            sealed_final_confirmation=args.confirmation,
        )
    else:
        result = verify_v4_artifacts(
            protocol_path=args.protocol,
            exploratory_path=args.exploratory,
            readiness_path=args.readiness,
            receipt_path=args.receipt,
            final_path=args.final,
        )
    print(json.dumps(result, sort_keys=True))
    if args.command in {"exploratory", "readiness"} and result["gate_passed"] is not True:
        raise SystemExit(1)
    if args.command == "sealed" and result["metrics"]["gate_passed"] is not True:
        raise SystemExit(1)
    if args.command == "verify" and result["verified"] is not True:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
